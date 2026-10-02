import asyncio
import io
from datetime import date

from mongomock_motor import AsyncMongoMockClient
from openpyxl import Workbook

from app.lotti.servizi.import_haccp_excel import prepara_importazione, sostituisci_archivio


def run(coro):
    return asyncio.run(coro)


def temperature_xlsx(nome, valori):
    wb = Workbook()
    ws = wb.active
    ws.title = nome
    ws["A2"] = "ANNO: 2026"
    ws["A5"] = "Giorno"
    for giorno in range(1, 32):
        ws.cell(5, giorno + 1, giorno)
    ws["A6"] = "GENNAIO"
    for giorno, valore in valori.items():
        ws.cell(6, giorno + 1, valore)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def sanificazione_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = "Gennaio"
    ws["A4"] = "Mese di: GENNAIO Anno 2026"
    ws["B6"] = "Attrezzature Laboratorio"
    ws["C6"] = "Utensili"
    ws["A9"] = 1
    ws["B9"] = "XX"
    ws["C9"] = "CHIUSI"
    ws["I9"] = "Vincenzo Ceraldi"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def preparata():
    return prepara_importazione(
        temperature_xlsx("CONGELATORE LAB", {1: "C", 2: -20}),
        "negative.xlsx",
        temperature_xlsx("FRIGO BAR", {1: "C", 2: 3}),
        "positive.xlsx",
        sanificazione_xlsx(),
        "sanificazione.xlsx",
        oggi=date(2026, 1, 2),
    )


def test_parser_preserva_valori_chiusura_nomi_e_sanificazioni():
    dati = preparata()
    neg = dati["negative"]
    pos = dati["positive"]
    san = dati["sanificazione"]

    assert neg["nomi"] == ["CONGELATORE LAB"]
    assert pos["nomi"] == ["FRIGO BAR"]
    assert neg["documenti"][0]["temperature"]["1"]["1"]["is_chiuso"] is True
    assert neg["documenti"][0]["temperature"]["1"]["2"]["temp"] == -20
    assert pos["documenti"][0]["temperature"]["1"]["2"]["temp"] == 3
    scheda = san["documenti"][0]
    assert scheda["registrazioni"]["Attrezzature Laboratorio"]["1"] == "X"
    assert scheda["valori_originali_excel"]["Attrezzature Laboratorio"]["1"] == "XX"
    assert scheda["registrazioni"]["Tagliere, Coltelli"]["1"] == "N/D"
    assert scheda["firme"]["Attrezzature Laboratorio"]["1"]["firma_verificata"] is False


def test_sostituzione_crea_backup_e_disattiva_solo_gli_apparecchi_eccedenti():
    db = AsyncMongoMockClient()["import_haccp_test"]
    run(db.temperature_positive.insert_one({"anno": 2026, "frigorifero_numero": 9, "temperature": {}}))
    run(db.attrezzature_config.insert_many([
        {"tipo": "frigo", "numero": 1, "nome": "Vecchio 1", "attivo": True},
        {"tipo": "frigo", "numero": 9, "nome": "Vecchio 9", "attivo": True},
    ]))

    esito = run(sostituisci_archivio(db, preparata(), {"nome": "Amministratore"}))

    assert esito["backup_id"]
    backup = run(db.haccp_import_backups.find_one({"id": esito["backup_id"]}))
    assert backup["stato"] == "applicato"
    assert backup["temperature_positive"][0]["frigorifero_numero"] == 9
    nuovo = run(db.temperature_positive.find_one({"anno": 2026, "frigorifero_numero": 1}))
    assert nuovo["frigorifero_nome"] == "FRIGO BAR"
    assert run(db.attrezzature_config.find_one({"tipo": "frigo", "numero": 1}))["nome"] == "FRIGO BAR"
    assert run(db.attrezzature_config.find_one({"tipo": "frigo", "numero": 9}))["attivo"] is False
