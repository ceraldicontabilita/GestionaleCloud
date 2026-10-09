import io
import asyncio

from openpyxl import Workbook
from starlette.datastructures import UploadFile

from app.hr.services.presenze_excel import (
    analizza_presenze_csv, analizza_presenze_workbook,
    indicizza_dipendenti_per_nome,
)


def _workbook_base():
    wb = Workbook()
    riepilogo = wb.active
    riepilogo.title = "Riepilogo Consulente"
    riepilogo.append(["Dipendente", "Codice fiscale", "Festività NON lavorate (pagate)"])
    giornaliere = wb.create_sheet("Presenze Giornaliere")
    giornaliere.append(["Dipendente", "Data", "Entrata", "Tipo timbratura", "Note"])
    assenze = wb.create_sheet("Assenze, Ferie, Malattie")
    assenze.append(["Dipendente", "Tipo", "Dal", "Al", "Protocollo INPS", "Note"])
    return wb, riepilogo, giornaliere, assenze


def test_assenza_prevale_sulla_timbratura_e_fnl_non_viene_inventata():
    wb, riepilogo, giornaliere, assenze = _workbook_base()
    riepilogo.append(["Mario Esempio", "RSSMRA93L21F839G", 1])
    giornaliere.append(["Mario Esempio", "01/09/2026", "08:00", "QR in cassa", None])
    giornaliere.append(["Mario Esempio", "02/09/2026", "08:05", "QR in cassa", "timbratura interna"])
    assenze.append(["Mario Esempio", "Malattia", "02/09/2026", "03/09/2026", "P123", "certificata"])

    risultato = analizza_presenze_workbook(wb)

    assert risultato["periodo"] == "2026-09"
    assert risultato["conteggi"] == {"P": 1, "M": 2, "RS": 27}
    assert "FNL" not in risultato["conteggi"]
    sovrapposta = next(r for r in risultato["record"] if r["data"] == "2026-09-02")
    assert sovrapposta["giustificativo"] == "M"
    assert sovrapposta["entrata"] == "08:05"
    assert "Protocollo INPS: P123" in sovrapposta["note"]


def test_nominativo_giornaliero_fuori_riepilogo_resta_associabile_in_anteprima():
    wb, riepilogo, giornaliere, _assenze = _workbook_base()
    riepilogo.append(["Francesco Iazzetta", "ZZTFNC07A05F839D", 0])
    giornaliere.append(["Liliana Strazzullo", "01/09/2026", "08:00", "QR in cassa", None])

    risultato = analizza_presenze_workbook(wb)

    assert risultato["conteggi"] == {"P": 1, "RS": 29}
    assert risultato["nominativi_da_associare"] == ["Liliana Strazzullo"]
    assert risultato["record"][0]["codice_fiscale"] is None
    assert risultato["errori"][0]["motivo"].startswith("nominativo giornaliero assente")


def test_griglia_xlsx_esportata_dalla_pagina_presenze():
    wb = Workbook()
    ws = wb.active
    ws.title = "presenze_2026_09"
    ws.append(["Presenze Settembre 2026 - Ceraldi Group S.r.l."])
    ws.append([])
    ws.append(["Dipendente", *range(1, 31)])
    ws.append(["Capezzuto Alessandro", None, "P", *([None] * 6), "M", *([None] * 21)])

    risultato = analizza_presenze_workbook(wb, "presenze_2026_09.xlsx")

    assert risultato["periodo"] == "2026-09"
    assert risultato["formato"] == "griglia"
    assert risultato["conteggi"] == {"P": 1, "M": 1}
    assert [r["data"] for r in risultato["record"]] == ["2026-09-02", "2026-09-09"]


def test_griglia_csv_semicolon_con_bom():
    righe = [
        "Presenze Settembre 2026 - Ceraldi Group S.r.l.",
        "",
        "Dipendente;" + ";".join(str(g) for g in range(1, 31)),
        "Lesina Angela;P;P;" + ";".join([""] * 28),
        "",
        "Legenda: P=Presente",
    ]
    risultato = analizza_presenze_csv(("\ufeff" + "\n".join(righe)).encode("utf-8"), "presenze_2026_09.csv")

    assert risultato["periodo"] == "2026-09"
    assert risultato["conteggi"] == {"P": 2}
    assert risultato["nominativi_da_associare"] == ["Lesina Angela"]


def test_indice_nomi_accetta_cognome_nome_e_nome_cognome_solo_se_univoci():
    mario = {"id": "1", "nome": "Mario", "cognome": "Rossi", "nome_completo": "Rossi Mario"}
    indice = indicizza_dipendenti_per_nome([mario])

    assert indice["ROSSI MARIO"] == [mario]
    assert indice["MARIO ROSSI"] == [mario]


class _Cursor:
    def __init__(self, righe):
        self.righe = righe

    async def to_list(self, _limite):
        return list(self.righe)


class _Collection:
    def __init__(self, righe):
        self.righe = righe

    def find(self, *_args, **_kwargs):
        return _Cursor(self.righe)


def test_endpoint_anteprima_accetta_csv_e_associa_nome_univoco(monkeypatch):
    from app.hr.routers import dipendenti_cloud as router

    db = type("DB", (), {})()
    db.dipendenti = _Collection([{
        "id": "dip-1", "nome": "Angela", "cognome": "Lesina",
        "nome_completo": "Lesina Angela", "codice_fiscale": "",
        "stato": "attivo", "attivo": True,
    }])
    db.presenze_cloud = _Collection([])
    monkeypatch.setattr(router, "get_db", lambda: db)
    contenuto = (
        "Presenze Settembre 2026 - Ceraldi Group S.r.l.\n\n"
        "Dipendente;" + ";".join(str(g) for g in range(1, 31)) + "\n"
        "Lesina Angela;P;P;" + ";".join([""] * 28)
    ).encode("utf-8")
    file = UploadFile(io.BytesIO(contenuto), filename="presenze_2026_09.csv")

    risultato = asyncio.run(router.importa_presenze_excel(
        file=file, applica=False, conferma_hash=None,
        sempre_presenti_ids="[]", sostituzioni_nomi="{}", sostituisci_mese=True,
    ))

    assert risultato["periodo"] == "2026-09"
    assert risultato["conteggi"]["nuova"] == 2
    assert risultato["nominativi_da_associare"] == []
    assert risultato["da_verificare"] == []
    assert risultato["associazioni_nomi"][0]["dipendente_id"] == "dip-1"
