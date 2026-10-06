"""Piano tributi: il piano crea l'attesa, le prove la soddisfano soltanto."""
import asyncio
from datetime import date

from mongomock_motor import AsyncMongoMockClient

from app.services import piano_tributi as piano
from app.services.expectation_policy import ExpectationStatus


def run(coro):
    return asyncio.run(coro)


def _f24(id_, versamento, erario=(), inps=(), regioni=(), **kw):
    return {
        "id": id_, "status": "da_pagare",
        "dati_generali": {"data_versamento": versamento, "saldo_delega_cents": kw.pop("saldo", 10000)},
        "sezione_erario": [
            {"codice_tributo": c, "anno": str(a), "mese": f"{m:02d}", "importo_debito_cents": d,
             "importo_credito_cents": cr, "periodo_riferimento": f"{m:02d}/{a}"}
            for c, m, a, d, cr in erario
        ],
        "sezione_inps": [
            {"causale": c, "anno": str(a), "mese": f"{m:02d}", "importo_debito_cents": d,
             "importo_credito_cents": 0} for c, m, a, d in inps
        ],
        "sezione_regioni": [
            {"codice_tributo": c, "anno": str(a), "mese": f"{m:02d}", "importo_debito_cents": d,
             "importo_credito_cents": 0} for c, m, a, d in regioni
        ],
        **kw,
    }


def _db(*f24):
    db = AsyncMongoMockClient()["piano"]
    if f24:
        run(db["f24_unificato"].insert_many([dict(f) for f in f24]))
    return db


def _voce(g, voce_id):
    return next(r for r in g["voci"] if r["voce"]["id"] == voce_id)


def _mese(g, voce_id, mese):
    return next(c for c in _voce(g, voce_id)["caselle"] if c["periodo"] == f"{mese:02d}")


OGGI = date(2026, 9, 26)


def test_il_piano_crea_l_attesa_prima_della_prova():
    g = run(piano.griglia(_db(), 2026, oggi=OGGI))
    giugno = _mese(g, "ritenute_1001", 6)
    assert giugno["stato"] == piano.MANCA_F24 and giugno["importo"] is None   # mai un importo inventato
    assert giugno["expectation_status"] == ExpectationStatus.ATTESO.value
    assert giugno["source_fact_id"] == "piano:ritenute_1001:2026:06"
    assert giugno["scadenza"] == "2026-07-16"
    assert _mese(g, "ritenute_1001", 10)["stato"] == piano.FUTURO
    assert {"anno": 2026, "voce": "Ritenute lavoro dipendente", "codici": ["1001"], "periodo": "06/2026",
            "scadenza": "2026-07-16", "stato": piano.MANCA_F24,
            "giorni_scaduto": (OGGI - date(2026, 7, 16)).days} in g["mancano"]


def test_la_prova_bancaria_soddisfa_e_conserva_gli_id():
    f = _f24("f-mag", "2026-06-16", erario=[("1001", 5, 2026, 102651, 0)],
             movimento_bancario_id="mov-1", data_pagamento_effettivo="2026-06-17")
    g = run(piano.griglia(_db(f), 2026, oggi=OGGI))
    maggio = _mese(g, "ritenute_1001", 5)
    assert maggio["stato"] == piano.PAGATO
    assert maggio["expectation_status"] == ExpectationStatus.SODDISFATTO.value
    assert maggio["importo"] == "1026.51"
    assert maggio["modelli"][0]["f24_id"] == "f-mag"


def test_la_sola_quietanza_non_vale_come_banca():
    f = _f24("f-apr", "2026-05-18", erario=[("1001", 4, 2026, 94079, 0)], quietanza_id="q-1")
    g = run(piano.griglia(_db(f), 2026, oggi=OGGI))
    aprile = _mese(g, "ritenute_1001", 4)
    assert aprile["stato"] == piano.QUIETANZA_SENZA_BANCA
    assert aprile["expectation_status"] == ExpectationStatus.DA_VERIFICARE.value


def test_f24_arrivato_e_non_pagato_dopo_la_scadenza():
    f = _f24("f-lug", "2026-08-20", erario=[("1001", 7, 2026, 179456, 0)])
    g = run(piano.griglia(_db(f), 2026, oggi=OGGI))
    assert _mese(g, "ritenute_1001", 7)["stato"] == piano.SCADUTO_NON_PAGATO
    assert _mese(g, "ritenute_1001", 7)["scadenza"] == "2026-08-20"   # 16 agosto -> 20
    g = run(piano.griglia(_db(f), 2026, oggi=date(2026, 8, 1)))
    assert _mese(g, "ritenute_1001", 7)["stato"] == piano.DA_PAGARE


def test_modello_doppio_conta_una_volta_e_vince_la_prova():
    pagato = _f24("f-a", "2026-05-18", erario=[("1001", 4, 2026, 94079, 0)], movimento_bancario_id="m", data_pagamento_effettivo="2026-06-17")
    doppio = _f24("f-b", "2026-05-18", erario=[("1001", 4, 2026, 94079, 0)])
    g = run(piano.griglia(_db(pagato, doppio), 2026, oggi=OGGI))
    aprile = _mese(g, "ritenute_1001", 4)
    assert aprile["stato"] == piano.PAGATO
    assert aprile["importo"] == "940.79"          # non raddoppiato
    assert aprile["modelli_doppi"] == 1 and g["modelli_doppi"] == 1


def test_addizionale_regionale_legge_le_rate_dell_anno_prima():
    f = _f24("f-mag", "2026-06-16", regioni=[("3802", 5, 2025, 69044)], movimento_bancario_id="m", data_pagamento_effettivo="2026-06-17")
    g = run(piano.griglia(_db(f), 2026, oggi=OGGI))
    assert _mese(g, "add_regionale_3802", 5)["stato"] == piano.PAGATO
    assert [c["periodo"] for c in _voce(g, "add_regionale_3802")["caselle"]][-1] == "11"


def test_inps_dm10_e_credito_mai_in_rosso():
    f = _f24("f-mag", "2026-06-16", inps=[("DM10", 5, 2026, 523700)],
             erario=[("1704", 5, 2026, 0, 67028)], movimento_bancario_id="m", data_pagamento_effettivo="2026-06-17")
    g = run(piano.griglia(_db(f), 2026, oggi=OGGI))
    assert _mese(g, "inps_dm10", 5)["stato"] == piano.PAGATO
    assert _mese(g, "credito_1704", 5)["stato"] == piano.CREDITO_USATO
    assert _mese(g, "credito_1704", 5)["credito"] == "670.28"
    assert _mese(g, "credito_1704", 6)["stato"] == piano.CREDITO_ASSENTE
    assert not _mese(g, "credito_1704", 6)["mandatory"]


def test_ires_saldo_dell_anno_prima_e_acconti():
    f = _f24("f-ires", "2026-06-30", erario=[("2003", 0, 2025, 500000, 0), ("2001", 0, 2026, 200000, 0)],
             movimento_bancario_id="m", data_pagamento_effettivo="2026-06-17")
    g = run(piano.griglia(_db(f), 2026, oggi=OGGI))
    assert _voce(g, "ires_saldo")["caselle"][0]["stato"] == piano.PAGATO
    assert _voce(g, "ires_acconto_1")["caselle"][0]["stato"] == piano.PAGATO
    assert _voce(g, "ires_acconto_2")["caselle"][0]["stato"] == piano.FUTURO
    assert _voce(g, "ires_acconto_2")["caselle"][0]["scadenza"] == "2026-11-30"


def test_cosap_fuori_f24_e_codici_fuori_piano():
    f = _f24("f-x", "2026-06-16", erario=[("1631", 5, 2026, 0, 1000)])
    g = run(piano.griglia(_db(f), 2026, oggi=OGGI))
    assert _voce(g, "cosap")["caselle"] == []                  # scadenza da impostare
    assert [v["codice"] for v in g["fuori_piano"]] == ["1631"]


def test_l_anno_non_si_chiude_con_un_attesa_aperta():
    g = run(piano.griglia(_db(), 2026, oggi=OGGI))
    assert g["anno_chiuso"] is False


def test_il_piano_si_scrive_una_volta_sola_e_rispetta_le_modifiche():
    db = _db()
    run(piano.griglia(db, 2026, oggi=OGGI))
    run(piano.aggiorna_voce(db, "inps_cxx", {"attivo": False}))
    run(db[piano.COLL_PIANO].delete_one({"id": "imu_saldo"}))    # voce di base nuova nel codice
    g = run(piano.griglia(db, 2026, oggi=OGGI))
    assert run(db[piano.COLL_PIANO].count_documents({})) == len(piano.PIANO_BASE)
    assert "inps_cxx" not in [r["voce"]["id"] for r in g["voci"]]      # spenta resta spenta
    assert "imu_saldo" in [r["voce"]["id"] for r in g["voci"]]


def test_modifica_voce_e_scadenza_cosap():
    db = _db()
    voce = run(piano.aggiorna_voce(db, "cosap", {"scadenze": [[3, 31]], "codici": ["X"]}))
    assert voce["scadenze"] == [[3, 31]] and "X" not in (voce.get("codici") or [])
    g = run(piano.griglia(db, 2026, oggi=OGGI))
    cosap = _voce(g, "cosap")["caselle"][0]
    assert cosap["stato"] == piano.DA_VERIFICARE_A_MANO and cosap["scadenza"] == "2026-03-31"


def test_scadenze_festivi_e_weekend():
    assert piano.scadenza_effettiva(date(2026, 8, 16)) == date(2026, 8, 20)
    assert piano.scadenza_effettiva(date(2026, 5, 16)) == date(2026, 5, 18)   # sabato
    assert piano.scadenza_effettiva(date(2026, 11, 1)) == date(2026, 11, 2)   # Ognissanti
    assert piano._scadenza_mensile(2026, 12) == date(2027, 1, 18)             # 16/01/2027 e' sabato


def test_rotte_solo_admin_e_prima_della_rotta_dinamica():
    from app.routers.f24.f24_main import router
    from app.utils.dependencies import get_current_admin_user

    percorsi = [r.path for r in router.routes]
    dinamica_get = next(i for i, r in enumerate(router.routes)
                        if r.path == "/{f24_id}" and "GET" in r.methods)
    for p in ("/piano-tributi", "/piano-tributi/voci", "/piano-tributi/voci/{voce_id}"):
        assert percorsi.index(p) < dinamica_get
        for rotta in (r for r in router.routes if r.path == p):
            assert get_current_admin_user in [d.call for d in rotta.dependant.dependencies]


def test_scadenza_festiva_slitta_e_pagare_il_primo_lavorativo_non_e_ritardo():
    """Il 16/05/2026 e' sabato: si paga lunedi' 18 senza ravvedimento (0 giorni di ritardo)."""
    f = _f24("f-apr", "2026-05-18", erario=[("1001", 4, 2026, 94079, 0)], movimento_bancario_id="m1")
    aprile = _mese(run(piano.griglia(_db(f), 2026, oggi=OGGI)), "ritenute_1001", 4)
    assert aprile["scadenza_nominale"] == "2026-05-16"
    assert aprile["scadenza"] == "2026-05-18"
    assert aprile["slittamento"] == "giorno festivo"
    assert aprile["giorni_ritardo"] == 0
    assert aprile["modelli"][0]["giorni_dopo_nominale"] == 2


def test_il_ritardo_si_conta_dal_termine_effettivo_e_il_ferragosto_va_al_20():
    ritardo = _f24("f-mag", "2026-06-19", erario=[("1001", 5, 2026, 102651, 0)], movimento_bancario_id="m2")
    maggio = _mese(run(piano.griglia(_db(ritardo), 2026, oggi=OGGI)), "ritenute_1001", 5)
    assert maggio["scadenza"] == "2026-06-16" and maggio["giorni_ritardo"] == 3
    # luglio 2026: il 16/08 slitta al 20/08 (proroga estiva), non e' un giorno festivo qualunque
    luglio = _mese(run(piano.griglia(_db(), 2026, oggi=OGGI)), "ritenute_1001", 7)
    assert luglio["scadenza_nominale"] == "2026-08-16" and luglio["scadenza"] == "2026-08-20"
    assert luglio["slittamento"] == "proroga di agosto"


def test_un_tributo_non_pagato_dice_da_quanti_giorni_e_scaduto():
    g = run(piano.griglia(_db(), 2026, oggi=OGGI))
    giugno = _mese(g, "ritenute_1001", 6)
    assert giugno["giorni_ritardo"] is None
    assert giugno["giorni_scaduto"] == (OGGI - date(2026, 7, 16)).days
    assert _mese(g, "ritenute_1001", 10)["giorni_scaduto"] is None   # non ancora scaduto


def test_piu_anni_e_tutti_dal_piu_recente():
    f = _f24("f-old", "2023-02-16", erario=[("1001", 1, 2023, 50000, 0)], movimento_bancario_id="m3")
    db = _db(f)
    g = run(piano.griglia_anni(db, "2024-2026", oggi=OGGI))
    assert g["multi"] is True and [x["anno"] for x in g["anni"]] == [2026, 2025, 2024]
    tutti = run(piano.griglia_anni(db, "tutti", oggi=OGGI))
    assert [x["anno"] for x in tutti["anni"]] == [2026, 2025, 2024, 2023]   # dal primo anno con un F24
    assert run(piano.griglia_anni(db, "2026", oggi=OGGI))["anno"] == 2026   # un anno solo: forma di sempre
    for sbagliato in ("abc", "2010", "2026-3000"):
        try:
            run(piano.griglia_anni(db, sbagliato, oggi=OGGI))
        except ValueError:
            continue
        raise AssertionError(f"{sbagliato!r} doveva essere rifiutato")


def test_codice_tributo_scritto_una_cifra_per_casella_si_riunisce():
    """Il modello del consulente (PCL2PDF) dà «1 0 0 1» in quattro caselle: va letto come 1001."""
    from app.services.parser_f24 import _unisci_codici_a_caselle

    def parola(x, w):
        return {"x": x, "x1": x + 6, "y": 225, "word": w}

    riga = [parola(161, "1"), parola(175, "0"), parola(190, "0"), parola(204, "1"),
            {"x": 233, "x1": 260, "y": 225, "word": "0002"}, {"x": 361, "x1": 390, "y": 225, "word": "384,66"}]
    assert [r["word"] for r in _unisci_codici_a_caselle(riga)] == ["1001", "0002", "384,66"]
    # un codice fiscale (11 cifre) o la regione «0 5» non si toccano
    cf = [parola(100 + 12 * i, str(i % 10)) for i in range(11)]
    assert len(_unisci_codici_a_caselle(cf)) == 11
    regione = [parola(24, "0"), parola(39, "5")]
    assert [r["word"] for r in _unisci_codici_a_caselle(regione)] == ["0", "5"]


def test_codice_comune_a_caselle_non_e_una_regione():
    """«E 9 0 6» (IMU) non e' la regione 06: il 3848 non finisce anche fra le regioni."""
    from app.services.parser_f24 import _codice_regione_da_riga

    def parola(x, w):
        return {"x": x, "x1": x + 6, "y": 520, "word": w}

    comune = [parola(25, "E"), parola(35, "9"), parola(45, "0"), parola(55, "6"), parola(170, "3848")]
    assert _codice_regione_da_riga(comune) == ""
    regione = [parola(30, "0"), parola(44, "5"), parola(170, "3802")]
    assert _codice_regione_da_riga(regione) == "05"


def test_excel_dello_scadenzario_un_versamento_per_riga_con_link_ritardo_e_prospetto():
    import io
    import openpyxl
    from app.services import scadenzario_excel as xl

    f1 = _f24("f-mag", "2026-06-19", erario=[("1001", 5, 2026, 102651, 0)], movimento_bancario_id="m1")
    f2 = _f24("f-gen", "2026-02-16", erario=[("1001", 1, 2026, 50000, 0), ("1655", 1, 2026, 0, 12000)])
    db = _db(f1, f2)
    run(db["prospetti_contabili"].insert_one({
        "id": "p1", "stato": "canonica", "f24_id": "f-gen", "mese": 1, "anno": 2026, "esito": "COMPLETO",
        "documento_id": "doc-1"}))
    righe = run(xl.righe_scadenzario(db, [2026]))
    assert [r["f24_id"] for r in righe] == ["f-mag", "f-gen", "f-gen"]          # il versamento piu' recente per primo
    assert righe[0]["giorni_ritardo"] == 3 and righe[1]["giorni_ritardo"] == 0   # 16/06 -> 19/06; 16/02 in termini
    wb = openpyxl.load_workbook(io.BytesIO(xl.costruisci_xlsx(righe, "https://gestionale.test")))
    assert wb.sheetnames == ["Scadenzario", "Riepilogo", "Debito e credito"]
    ws = wb["Scadenzario"]
    assert [c.value for c in ws[1]][:8] == [
        "Data", "Descrizione", "rateazione, regione/provincia, mese rif.", "anno di riferimento",
        "Codice tributo", "Importo", "f24 consulente", "quietanza"]
    assert ws["E2"].value == 1001 and ws["F2"].value == 1026.51 and ws["J2"].value == 3
    assert ws["G2"].hyperlink.target == "https://gestionale.test/fiscale/f24/f-mag"
    assert ws["K4"].value == "01/2026 · COMPLETO" and "doc-1" in ws["K4"].hyperlink.target
    assert ws["F4"].value == -120.0                                              # un credito compensato e' negativo
    riepilogo = {(r[0].value, r[1].value): r[3].value for r in wb["Riepilogo"].iter_rows(min_row=2) if r[1].value}
    assert riepilogo[(2026, 1001)] == 500.0                                      # febbraio: 16/02, 1001 di gennaio
    assert wb["Debito e credito"]["G4"].value == 120.0
    assert run(xl.righe_scadenzario(db, [2024])) == []                           # un anno senza versamenti: vuoto


def test_i_documenti_f24_sono_ordinati_per_data_dal_piu_recente_e_senza_data_in_fondo():
    from app.services.registro_fiscale_f24 import documenti_f24

    def riga(doc, data, nome):
        return {"document_id": doc, "payment_date": data, "filename": nome, "debit_amount": 1, "credit_amount": 0}

    righe = [riga("a", "2026-01-16", "a.pdf"), riga("b", None, "b.pdf"),
             riga("c", "2026-08-20", "c.pdf"), riga("d", "2026-03-16", "d.pdf")]
    assert [d["filename"] for d in documenti_f24(righe)] == ["c.pdf", "d.pdf", "a.pdf", "b.pdf"]


def test_modello_superato_o_in_quarantena_non_e_un_modello_del_registro():
    """«Versione superata» lascia lo stato `da_pagare`: la vecchia stampa non deve far risultare il mese non pagato."""
    from app.services.f24_controllo_incrociato import modello_attivo

    assert modello_attivo({"id": "a", "status": "da_pagare"})
    assert not modello_attivo({"id": "b", "status": "da_pagare", "superato_da": "a", "motivo_quarantena": "versione_superata"})
    assert not modello_attivo({"id": "c", "status": "da_pagare", "motivo_quarantena": "doppione"})
    assert not modello_attivo({"id": "d", "entity_status": "deleted"})


def test_gestione_separata_non_e_obbligatoria_nessun_f24_non_e_manca_f24():
    """CXX si versa solo se nel mese ci sono compensi: l'assenza dell'F24 non e' una mancanza dichiarata."""
    g = run(piano.griglia(_db(), 2026, oggi=OGGI))
    gennaio = _mese(g, "inps_cxx", 1)
    assert gennaio["stato"] == piano.NESSUN_F24 and gennaio["etichetta_stato"] == "Nessun F24 in archivio"
    assert gennaio["giorni_scaduto"] is None and gennaio["mandatory"] is False
    assert not any(m["codici"] == ["CXX"] for m in g["mancano"])
    # le voci obbligatorie restano «Manca F24»
    assert _mese(g, "ritenute_1001", 6)["stato"] == piano.MANCA_F24
