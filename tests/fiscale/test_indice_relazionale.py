"""MINI-06: indice relazionale ed export, e le relazioni documentali (DRV-03 minimo).

Si prova che: la relazione entita' -> documento nasce una volta sola (il secondo
giro non crea niente), il dry_run non scrive, l'ambiguo va in DA_VERIFICARE e
non si indovina, una relazione revocata non rinasce, l'esportazione e'
riproducibile (stessi byte) e i golden non contengono dati personali.
"""
import asyncio
import csv
import io
import json
import re
import zipfile
from pathlib import Path

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from app.services import indice_relazionale as indice
from app.services import relazioni_documentali as doc
from app.services.archivio_documenti_memoria import ClientArchivioMemoria
from app.services.entity_relations import upsert_entity_relation

GOLDEN = Path(__file__).parent / "golden_indice_relazionale"
MD5_A = "a" * 32
MD5_B = "b" * 32


def _run(coro):
    loop = asyncio.new_event_loop()
    try:
        return loop.run_until_complete(coro)
    finally:
        loop.close()


def _db(nome):
    return ClientArchivioMemoria()[nome]


def _ins(db, coll, **campi):
    # nel runtime la chiave e' `_id` (regola 10): `id` e' un campo dei dati, a volte numerico
    if "id" in campi:
        campi["_id"] = str(campi["id"])
        if coll != "invoices":
            del campi["id"]
    _run(db[coll].insert_one(campi))


def _relazioni(db):
    return _run(db["entity_relations"].find({}, {"_id": 0}).to_list(length=None))


def _archivio():
    db = _db("indice-doc")
    _ins(db, "invoices", id=1, drive_file_id="D-INV-1", status="imported")
    _ins(db, "invoices", id=2, drive_file_id="D-INV-1", status="archived")      # copia archiviata: fuori
    _ins(db, "invoices", id=3, drive_file_id="D-INV-3", status="imported")
    _ins(db, "f24_unificato", id="f24-1", drive_file_id="D-F24-1")
    _ins(db, "f24_unificato", id="f24-2", drive_file_id="D-F24-2", status="eliminato")  # quarantena: fuori
    _ins(db, "quietanze_f24", id="q-1", drive_file_id="D-Q-1")
    _ins(db, "invoices", id=4, status="imported")                                # senza file: nessuna relazione
    return db


PROTOCOLLO = [
    {"drive_id": "D-CED-1", "md5": MD5_A, "collegamento_tipo": "hr_cedolino", "collegamento_id": "hr-1"},
    {"drive_id": "D-Q-1", "md5": "c" * 32, "collegamento_tipo": "quietanze_f24", "collegamento_id": "q-1"},
    {"drive_id": "D-X", "md5": "d" * 32, "collegamento_tipo": "tipo_sconosciuto", "collegamento_id": "z"},
    {"drive_id": "D-SENZA", "md5": "e" * 32, "collegamento_tipo": None, "collegamento_id": None},
]


# ── relazioni documentali ────────────────────────────────────────────────────

def test_dry_run_non_scrive_e_conta():
    db = _archivio()
    esito = _run(doc.esegui(db, protocollo=PROTOCOLLO))
    assert esito["dry_run"] is True and esito["scritte"] == 0
    assert _relazioni(db) == []
    # invoice 1 e 3, f24-1, q-1 (campo + protocollo = una sola), hr-1
    assert esito["previste"] == 5 and esito["nuove"] == 5 and esito["da_verificare"] == 0
    assert esito["collegamenti_non_mappati"] == {"tipo_sconosciuto": 1}


def test_relazione_creata_una_volta_e_secondo_giro_zero_nuove():
    db = _archivio()
    primo = _run(doc.esegui(db, dry_run=False, protocollo=PROTOCOLLO))
    assert primo["scritte"] == 5
    relazioni = _relazioni(db)
    assert len(relazioni) == 5
    quietanza = next(r for r in relazioni if r["source"]["id"] == "q-1")
    assert quietanza["target"] == {"type": "documento", "id": "D-Q-1"}
    assert quietanza["relation_type"] == "has_source_document" and quietanza["status"] == "confirmed"
    assert quietanza["rule"] == "drive_file_id" and quietanza["provenance"]["fonti"] == ["campo", "protocollo"]
    creata = quietanza["created_at"]

    secondo = _run(doc.esegui(db, dry_run=False, protocollo=PROTOCOLLO))
    assert secondo["nuove"] == 0 and secondo["aggiornate"] == 0 and secondo["scritte"] == 0
    assert secondo["gia_presenti"] == 5
    assert len(_relazioni(db)) == 5
    assert next(r for r in _relazioni(db) if r["source"]["id"] == "q-1")["created_at"] == creata


def test_archiviata_e_quarantena_non_generano_relazioni():
    db = _archivio()
    _run(doc.esegui(db, dry_run=False, protocollo=[]))
    chiavi = {(r["source"]["type"], r["source"]["id"]) for r in _relazioni(db)}
    assert ("invoice", "2") not in chiavi and ("f24_model", "f24-2") not in chiavi
    assert ("invoice", "4") not in chiavi


def test_entita_con_due_file_diversi_e_da_verificare():
    db = _archivio()
    protocollo = [{"drive_id": "D-ALTRO", "md5": "f" * 32, "collegamento_tipo": "quietanze_f24", "collegamento_id": "q-1"}]
    esito = _run(doc.esegui(db, dry_run=False, protocollo=protocollo))
    q = [r for r in _relazioni(db) if r["source"]["id"] == "q-1"]
    assert {r["target"]["id"] for r in q} == {"D-Q-1", "D-ALTRO"}
    assert {r["status"] for r in q} == {"pending"}
    assert {r["provenance"]["motivo"] for r in q} == {doc.MOTIVO_PIU_FILE}
    assert esito["da_verificare"] == 2
    assert {a["entity_id"] for a in esito["ambigue"]} == {"q-1"}


def test_file_con_due_fatture_attive_e_da_verificare_ma_non_i_cedolini():
    db = _db("indice-doppia")
    _ins(db, "invoices", id=10, drive_file_id="D-DUE", status="imported")
    _ins(db, "invoices", id=11, drive_file_id="D-DUE", status="imported")
    protocollo = [
        {"drive_id": "D-COMB", "md5": MD5_A, "collegamento_tipo": "hr_cedolino", "collegamento_id": "hr-1"},
        {"drive_id": "D-COMB", "md5": MD5_A, "collegamento_tipo": "hr_cedolino", "collegamento_id": "hr-2"},
    ]
    esito = _run(doc.esegui(db, dry_run=False, protocollo=protocollo))
    per_id = {r["source"]["id"]: r["status"] for r in _relazioni(db)}
    assert per_id["10"] == per_id["11"] == "pending"
    assert per_id["hr-1"] == per_id["hr-2"] == "confirmed"      # documento combinato: non e' ambiguita'
    assert esito["da_verificare"] == 2


def test_cedolino_per_impronta_md5_nota_e_impronta_sha256_fuori():
    db = _db("indice-ced")
    _ins(db, "cedolini", id="c1", source_file_hash=MD5_A)
    _ins(db, "cedolini", id="c2", source_file_hash=MD5_B)               # nessun file con questa impronta
    _ins(db, "cedolini", id="c3", source_file_hash="9" * 64)            # SHA-256: il protocollo non ce l'ha
    _ins(db, "cedolini", id="c4", source_file_hash=MD5_A, status="sostituito")
    protocollo = [{"drive_id": "D-CED", "md5": MD5_A, "collegamento_tipo": None, "collegamento_id": None}]
    esito = _run(doc.esegui(db, dry_run=False, protocollo=protocollo))
    assert [(r["source"]["type"], r["source"]["id"], r["target"]["id"]) for r in _relazioni(db)] == [("payslip", "c1", "D-CED")]
    busta = esito["per_fonte_dati"]["payslip"]
    assert busta["con_impronta"] == 3 and busta["con_file"] == 1
    assert busta["impronta_senza_file"] == 2        # c2 (MD5 sconosciuta) e c3 (SHA-256 fuori dal registro)


def test_impronta_su_due_file_e_ambigua():
    db = _db("indice-ced2")
    _ins(db, "cedolini", id="c1", source_file_hash=MD5_A)
    protocollo = [{"drive_id": "D1", "md5": MD5_A}, {"drive_id": "D2", "md5": MD5_A}]
    _run(doc.esegui(db, dry_run=False, protocollo=protocollo))
    assert {r["status"] for r in _relazioni(db)} == {"pending"} and len(_relazioni(db)) == 2


def test_relazione_revocata_non_rinasce_e_pending_si_promuove():
    db = _archivio()
    _run(doc.esegui(db, dry_run=False, protocollo=[]))
    _run(db["entity_relations"].update_one(
        {"relation_key": "invoice|1|has_source_document|documento|D-INV-1"}, {"$set": {"status": "revoked"}}))
    _run(db["entity_relations"].update_one(
        {"relation_key": "invoice|3|has_source_document|documento|D-INV-3"}, {"$set": {"status": "pending"}}))
    esito = _run(doc.esegui(db, dry_run=False, protocollo=[]))
    assert esito["revocate_rispettate"] == 1 and esito["aggiornate"] == 1 and esito["nuove"] == 0
    stati = {r["source"]["id"]: r["status"] for r in _relazioni(db) if r["source"]["type"] == "invoice"}
    assert stati == {"1": "revoked", "3": "confirmed"}


def test_collegamento_del_protocollo_verso_entita_inesistente_o_fuori_si_scarta():
    db = _archivio()
    protocollo = [
        {"drive_id": "D-FANTASMA", "md5": "1" * 32, "collegamento_tipo": "quietanze_f24", "collegamento_id": "q-non-esiste"},
        {"drive_id": "D-ARCH", "md5": "2" * 32, "collegamento_tipo": "f24_unificato", "collegamento_id": "f24-2"},  # quarantena
    ]
    esito = _run(doc.esegui(db, dry_run=False, protocollo=protocollo))
    assert esito["collegamenti_non_mappati"] == {"quietanze_f24:entita_assente": 1, "f24_unificato:entita_assente": 1}
    assert not {r["target"]["id"] for r in _relazioni(db)} & {"D-FANTASMA", "D-ARCH"}


def test_la_chiave_e_l_id_del_runtime_anche_se_la_fattura_ha_un_id_numerico():
    db = _db("indice-chiave")
    _run(db["invoices"].insert_one({"_id": "uuid-fattura", "id": 77, "drive_file_id": "D-77", "status": "imported"}))
    protocollo = [{"drive_id": "D-77", "md5": "3" * 32, "collegamento_tipo": "invoice", "collegamento_id": "uuid-fattura"}]
    esito = _run(doc.esegui(db, dry_run=False, protocollo=protocollo))
    rel = _relazioni(db)
    assert [r["source"]["id"] for r in rel] == ["uuid-fattura"] and esito["previste"] == 1   # campo e protocollo: una relazione


def test_protocollo_non_raggiungibile_si_dichiara_non_si_inventa():
    db = _archivio()

    async def _guasto():
        raise RuntimeError("DSN")

    esito = _run(doc.esegui(db, leggi_prot=_guasto))
    assert esito["protocollo"] == "non_disponibile" and esito["impronte_md5"] == "protocollo_non_disponibile"
    assert esito["previste"] == 4         # solo i campi drive_file_id


# ── indice ed export ─────────────────────────────────────────────────────────

def _fixture_indice():
    db = _db("indice-export")

    async def _crea():
        await upsert_entity_relation(
            db, source_type="invoice", source_id="INV-A", relation_type="has_source_document",
            target_type="documento", target_id="DRIVE-001", status="confirmed", rule="drive_file_id",
            evidence=[{"type": "drive_file_id", "value": "DRIVE-001"}, {"type": "md5", "value": MD5_A}])
        await upsert_entity_relation(
            db, source_type="bank_movement", source_id="MOV-7", relation_type="proves_invoice_payment",
            target_type="invoice", target_id="INV-A", status="pending", rule="cro",
            evidence=[{"type": "cro", "value": "CRO-0001"}], amount="1234.56")
        await upsert_entity_relation(
            db, source_type="f24_receipt", source_id="=CMD", relation_type="has_source_document",
            target_type="documento", target_id="DRIVE-002", status="revoked", rule="collegamento_protocollo")
    _run(_crea())
    # date fisse: il golden non dipende dall'orologio
    for chiave, data in (("INV-A", "2026-09-01T10:00:00+00:00"), ("MOV-7", "2026-09-15T08:30:00+00:00"),
                         ("=CMD", "2026-09-30T23:59:59+00:00")):
        _run(db["entity_relations"].update_one({"source.id": chiave}, {"$set": {"created_at": data, "updated_at": data}}))
    return db


def test_export_e_riproducibile_stessi_byte_per_ogni_formato():
    db = _fixture_indice()
    for formato in indice.FORMATI:
        uno = indice.esporta(_run(indice.leggi_righe(db))["righe"], formato)
        due = indice.esporta(_run(indice.leggi_righe(db))["righe"], formato)
        assert uno == due, formato
        assert uno[2] == f"indice_relazionale.{formato}"


def test_export_ordinato_per_chiave_e_importi_decimal_date_italiane():
    db = _fixture_indice()
    righe = _run(indice.leggi_righe(db))["righe"]
    assert [r["relation_key"] for r in righe] == sorted(r["relation_key"] for r in righe)
    corpo = indice.esporta_csv(righe).decode("utf-8-sig")
    lette = list(csv.reader(io.StringIO(corpo), delimiter=";"))
    assert lette[0] == [c for _, c in indice.COLONNE]
    mov = next(r for r in lette if r[2] == "MOV-7")
    assert mov[8] == "1234,56" and mov[10] == "15/09/2026" and mov[6] == "DA_VERIFICARE"
    assert mov[9] == "cro=CRO-0001"
    assert any(r[2] == "'=CMD" for r in lette)          # niente formula in Excel
    js = json.loads(indice.esporta_json(righe))
    mov_j = next(r for r in js["relazioni"] if r["origine_id"] == "MOV-7")
    assert mov_j["importo"] == "1234.56" and mov_j["creata"] == "2026-09-15T08:30:00+00:00"


def test_xlsx_ha_le_stesse_righe_del_csv_e_zip_a_data_fissa():
    from openpyxl import load_workbook
    db = _fixture_indice()
    righe = _run(indice.leggi_righe(db))["righe"]
    dati = indice.esporta_xlsx(righe)
    with zipfile.ZipFile(io.BytesIO(dati)) as z:
        assert {i.date_time for i in z.infolist()} == {(2000, 1, 1, 0, 0, 0)}
    ws = load_workbook(io.BytesIO(dati)).active
    valori = [[c.value for c in r] for r in ws.iter_rows()]
    assert valori[0] == [c for _, c in indice.COLONNE] and len(valori) == 1 + len(righe)
    mov = next(r for r in valori if r[2] == "MOV-7")
    assert float(mov[8]) == 1234.56 and mov[10] == "15/09/2026"


def test_filtri_e_riepilogo():
    db = _fixture_indice()
    solo_doc = _run(indice.leggi_righe(db, indice.costruisci_filtro(destinazione_tipo="documento")))["righe"]
    assert len(solo_doc) == 2
    da_ver = _run(indice.leggi_righe(db, indice.costruisci_filtro(stato="DA_VERIFICARE")))["righe"]
    assert [r["origine_id"] for r in da_ver] == ["MOV-7"]
    lato = _run(indice.leggi_righe(db, indice.costruisci_filtro(entita_id="INV-A")))["righe"]
    assert len(lato) == 2          # come origine e come destinazione
    tutte = _run(indice.leggi_righe(db))["righe"]
    assert indice.riepilogo(tutte)["per_stato"] == {"CONFERMATA": 1, "DA_VERIFICARE": 1, "REVOCATA": 1}


def test_formato_non_valido():
    with pytest.raises(ValueError):
        indice.esporta([], "pdf")


def test_golden_indice_uguale_all_export_e_senza_dati_personali():
    db = _fixture_indice()
    righe = _run(indice.leggi_righe(db))["righe"]
    assert indice.esporta_csv(righe) == (GOLDEN / "indice.csv").read_bytes().replace(b"\r\n", b"\n")
    assert indice.esporta_json(righe) == (GOLDEN / "indice.json").read_bytes().replace(b"\r\n", b"\n")
    testo = (GOLDEN / "indice.csv").read_text("utf-8") + (GOLDEN / "indice.json").read_text("utf-8")
    assert not re.search(r"\b[A-Z]{6}\d{2}[A-Z]\d{2}[A-Z]\d{3}[A-Z]\b", testo)       # codice fiscale
    assert not re.search(r"\bIT\d{2}[A-Z]\d{10}[A-Z0-9]{12}\b", testo)                # IBAN
    assert not re.search(r"[\w.]+@[\w.]+", testo)                                      # email


def test_golden_minisito_212_quietanze_una_relazione_ciascuna_secondo_giro_zero():
    """Le 212 quietanze del golden del minisito diventano 212 chiavi distinte."""
    manifest = json.loads((Path(__file__).parent / "golden_minisito" / "manifest_final.json").read_text("utf-8"))
    quietanze = manifest["raw"]["f24"]
    db = _db("indice-golden-minisito")
    for n, q in enumerate(quietanze):
        _ins(db, "quietanze_f24", id=f"gq-{n:03d}", drive_file_id=f"drive-{n:03d}")
    primo = _run(doc.esegui(db, dry_run=False, protocollo=[]))
    secondo = _run(doc.esegui(db, dry_run=False, protocollo=[]))
    assert primo["nuove"] == len(quietanze) == 212 and primo["da_verificare"] == 0
    assert secondo["nuove"] == 0 and secondo["scritte"] == 0
    assert len({r["relation_key"] for r in _relazioni(db)}) == 212


# ── endpoint ─────────────────────────────────────────────────────────────────

def _client(monkeypatch, db):
    from app.database import Database
    from app.routers import indice_relazionale as router
    from app.utils.dependencies import get_current_admin_user

    monkeypatch.setattr(Database, "get_db", staticmethod(lambda: db))
    monkeypatch.setattr(router, "_job_task", None)

    async def _senza_protocollo():
        return []

    monkeypatch.setattr(doc, "leggi_protocollo", _senza_protocollo)
    app = FastAPI()
    app.include_router(router.router, prefix="/api/indice-relazionale")
    app.dependency_overrides[get_current_admin_user] = lambda: {"sub": "a", "role": "admin"}
    return TestClient(app)


def test_endpoint_elenco_export_e_backfill(monkeypatch):
    db = _fixture_indice()
    _ins(db, "invoices", id=99, drive_file_id="D-99", status="imported")
    with _client(monkeypatch, db) as client:
        elenco = client.get("/api/indice-relazionale", params={"limit": 2}).json()
        assert elenco["riepilogo"]["totale"] == 3 and len(elenco["righe"]) == 2
        assert elenco["righe"][0]["aggiornata"] >= elenco["righe"][1]["aggiornata"]      # il piu' recente prima
        export = client.get("/api/indice-relazionale/export", params={"formato": "csv", "stato": "CONFERMATA"})
        assert export.status_code == 200 and "attachment" in export.headers["content-disposition"]
        assert export.content == client.get("/api/indice-relazionale/export", params={"formato": "csv", "stato": "CONFERMATA"}).content
        assert client.get("/api/indice-relazionale/export", params={"formato": "pdf"}).status_code == 400

        prova = client.post("/api/indice-relazionale/relazioni-documentali/backfill")
        assert prova.json()["avviato"] is True and prova.json()["dry_run"] is True
        for _ in range(100):
            stato = client.get("/api/indice-relazionale/relazioni-documentali/stato").json()
            if stato.get("stato") == "completato":
                break
        assert stato["esito"]["dry_run"] is True and stato["esito"]["nuove"] == 1
        vero = client.post("/api/indice-relazionale/relazioni-documentali/backfill", params={"dry_run": "false"})
        assert vero.json()["avviato"] is True
        for _ in range(100):
            stato = client.get("/api/indice-relazionale/relazioni-documentali/stato").json()
            if stato.get("stato") == "completato" and stato["esito"]["dry_run"] is False:
                break
        assert stato["esito"]["scritte"] == 1
    assert len(_run(db["entity_relations"].find({"source.id": "99"}).to_list(length=None))) == 1


def test_endpoint_solo_admin():
    from app.routers import indice_relazionale as router
    from app.utils.dependencies import get_current_admin_user
    for rotta in router.router.routes:
        assert any(d.call is get_current_admin_user for d in rotta.dependant.dependencies), rotta.path
