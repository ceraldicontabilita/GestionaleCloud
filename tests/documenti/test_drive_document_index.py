import io

import pytest
from openpyxl import Workbook

from app.services.drive_document_index import (
    DECLARATION_HEADERS,
    DUPLICATE_HEADERS,
    F24_HEADERS,
    FOLDER_MIME,
    REQUIRED_HEADERS,
    _discover_index_file_sync,
    _parse_index_workbook,
    validate_relations,
)


class _Call:
    def __init__(self, value):
        self.value = value

    def execute(self):
        return self.value


class _Files:
    def __init__(self, children, metadata):
        self.children = children
        self.metadata = metadata

    def list(self, q, **kwargs):
        parent = q.split("'")[1]
        return _Call({"files": self.children.get(parent, [])})

    def get(self, fileId, **kwargs):
        return _Call(self.metadata[fileId])


class _Service:
    def __init__(self, children, metadata):
        self._files = _Files(children, metadata)

    def files(self):
        return self._files


def _xlsx(records, f24_rows=None, declarations=None, duplicates=None):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "DOCUMENTI"
    sheet.append(REQUIRED_HEADERS)
    for record in records:
        sheet.append([record.get(header) for header in REQUIRED_HEADERS])
    for name, headers, values in (
        ("F24_RIGHE", F24_HEADERS, f24_rows or []),
        ("DICHIARAZIONI", DECLARATION_HEADERS, declarations or []),
        ("DUPLICATI_SCARTI", DUPLICATE_HEADERS, duplicates or []),
    ):
        child = workbook.create_sheet(name)
        child.append(headers)
        for record in values:
            child.append([record.get(header) for header in headers])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_discovers_one_exact_index_and_fails_closed_on_duplicates():
    root = {"id": "root", "name": "GESTIONALE", "mimeType": FOLDER_MIME, "trashed": False}
    folder = {"id": "indexes", "name": "INDICI GESTIONALE", "mimeType": FOLDER_MIME}
    index = {"id": "xlsx", "name": "INDICE_DOCUMENTALE_DRIVE.xlsx", "mimeType": "application/xlsx"}
    metadata = {"root": root, "xlsx": {**index, "trashed": False, "webViewLink": "https://drive/xlsx"}}
    service = _Service({"root": [folder], "indexes": [index]}, metadata)

    found_folder, found_index = _discover_index_file_sync(service, "root")
    assert found_folder["id"] == "indexes"
    assert found_index["id"] == "xlsx"

    duplicate_service = _Service({"root": [folder, {**folder, "id": "indexes-2"}]}, metadata)
    with pytest.raises(ValueError, match="ambiguo"):
        _discover_index_file_sync(duplicate_service, "root")


def test_rejects_workbook_without_required_headers():
    workbook = Workbook()
    workbook.active.title = "DOCUMENTI"
    workbook.active.append(["ID documento"])
    output = io.BytesIO()
    workbook.save(output)
    with pytest.raises(ValueError, match="Colonne obbligatorie assenti"):
        _parse_index_workbook(output.getvalue())


def test_boolean_validation_connects_f24_and_declaration_to_document():
    document = {
        "ID documento": "DOC-1", "Nome file": "modello.pdf", "SHA-256": "a" * 64,
        "Percorso Drive": r"F24\2026\modello.pdf",
    }
    catalog = {
        "documents": [document],
        "f24_rows": [{
            "ID documento": "DOC-1", "SHA-256": "a" * 64,
            "Percorso Drive": r"F24\2026\modello.pdf", "Debito": 10, "Credito": 0,
        }],
        "declarations": [{"Percorso archivio": "02_ANNI/2026/DICHIARAZIONI/modello.pdf"}],
        "duplicates": [],
    }
    validation = validate_relations(catalog)
    assert validation["all_true"] is True
    assert all(validation["checks"].values())
    assert validation["counts"]["f24_documents"] == 1


def test_boolean_validation_rejects_broken_relations():
    catalog = {
        "documents": [{
            "ID documento": "DOC-1", "Nome file": "modello.pdf",
            "SHA-256": "a" * 64, "Percorso Drive": r"F24\modello.pdf",
        }],
        "f24_rows": [{
            "ID documento": "DOC-MISSING", "SHA-256": "b" * 64,
            "Percorso Drive": r"F24\altro.pdf", "Debito": -1,
        }],
        "declarations": [{"Percorso archivio": "missing.pdf"}],
        "duplicates": [],
    }
    validation = validate_relations(catalog)
    assert validation["all_true"] is False
    assert validation["checks"]["all_f24_documents_exist"] is False
    assert validation["checks"]["all_declarations_link_exactly_one_document"] is False
