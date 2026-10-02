"""Importazione verificabile dei registri HACCP Excel 2026.

I file sono fonti documentali: i valori memorizzati da Excel vengono copiati
senza ricalcolare le formule volatili. I giorni futuri non vengono mai importati.
"""

from __future__ import annotations

import calendar
import hashlib
import io
import re
import uuid
from datetime import date, datetime, time, timezone
from typing import Any

from fastapi import HTTPException
from openpyxl import load_workbook

from app.lotti.servizi.registro_haccp import FUSO


MESI = {
    "GENNAIO": 1,
    "FEBBRAIO": 2,
    "MARZO": 3,
    "APRILE": 4,
    "MAGGIO": 5,
    "GIUGNO": 6,
    "LUGLIO": 7,
    "AGOSTO": 8,
    "SETTEMBRE": 9,
    "OTTOBRE": 10,
    "NOVEMBRE": 11,
    "DICEMBRE": 12,
}

SANIFICAZIONE_ALIASES = {
    "ATTREZZATURE LABORATORIO": "Attrezzature Laboratorio",
    "UTENSILI": "Tagliere, Coltelli",
    "PAVIMENTAZIONE": "Pavimentazione",
    "ATTREZZATURE BAR": "Attrezzature Bar",
    "MONTACARICHI": "Montacarichi",
    "DEPOSITO": "Deposito",
}


def _sha256(raw: bytes) -> str:
    return hashlib.sha256(raw).hexdigest().upper()


def _anno(wb) -> int:
    for ws in wb.worksheets:
        for coord in ("A2", "A4"):
            match = re.search(r"\b(20\d{2})\b", str(ws[coord].value or ""))
            if match:
                return int(match.group(1))
    raise HTTPException(422, "Anno non riconosciuto nel file Excel")


def _workbooks(raw: bytes):
    try:
        return (
            load_workbook(io.BytesIO(raw), data_only=False, read_only=False),
            load_workbook(io.BytesIO(raw), data_only=True, read_only=False),
        )
    except Exception as exc:
        raise HTTPException(422, f"File Excel non leggibile: {exc}") from exc


def _timestamp_giorno(giorno: date) -> str:
    locale = datetime.combine(giorno, time(hour=7), tzinfo=FUSO)
    return locale.astimezone(timezone.utc).isoformat()


def leggi_temperature(raw: bytes, filename: str, tipo: str, oggi: date | None = None) -> dict:
    """Legge un file a fogli: ogni foglio e' un apparecchio."""
    oggi = oggi or datetime.now(FUSO).date()
    wf, wv = _workbooks(raw)
    anno = _anno(wf)
    if anno != oggi.year:
        raise HTTPException(422, f"{filename}: anno {anno}, atteso {oggi.year}")
    if len(set(wf.sheetnames)) != len(wf.sheetnames):
        raise HTTPException(422, f"{filename}: nomi foglio duplicati")

    minimo, massimo = ((0.0, 4.0) if tipo == "frigo" else (-22.0, -18.0))
    campo_numero = "frigorifero_numero" if tipo == "frigo" else "congelatore_numero"
    campo_nome = "frigorifero_nome" if tipo == "frigo" else "congelatore_nome"
    docs = []
    totale = formule = chiusi = anomalie = 0
    mancanti: list[str] = []
    generici: list[str] = []
    digest = _sha256(raw)

    for numero, nome in enumerate(wf.sheetnames, 1):
        if re.fullmatch(r"Foglio\d+", nome, flags=re.IGNORECASE):
            generici.append(nome)
        sf, sv = wf[nome], wv[nome]
        temperature = {str(m): {} for m in range(1, 13)}
        date_viste: set[date] = set()
        for row in range(1, sf.max_row + 1):
            mese = MESI.get(str(sf.cell(row, 1).value or "").strip().upper())
            if not mese:
                continue
            for giorno in range(1, calendar.monthrange(anno, mese)[1] + 1):
                data = date(anno, mese, giorno)
                if data > oggi:
                    continue
                formula_o_valore = sf.cell(row, giorno + 1).value
                is_formula = isinstance(formula_o_valore, str) and formula_o_valore.startswith("=")
                valore = sv.cell(row, giorno + 1).value if is_formula else formula_o_valore
                if valore is None or str(valore).strip() == "":
                    mancanti.append(f"{nome}:{data.isoformat()}")
                    continue
                if data in date_viste:
                    raise HTTPException(422, f"{filename}: data duplicata {data} nel foglio {nome}")
                date_viste.add(data)
                base = {
                    "timestamp": _timestamp_giorno(data),
                    "firma_verificata": False,
                    "firma_via": "import_excel",
                    "fonte_importazione": filename,
                    "fonte_sha256": digest,
                }
                testo = str(valore).strip().upper()
                if testo in {"C", "CHIUSO", "CHIUSI"}:
                    record = {
                        **base,
                        "temp": None,
                        "stato": "chiuso",
                        "is_chiuso": True,
                        "allarme": False,
                        "valore_originale": str(valore),
                    }
                    chiusi += 1
                elif isinstance(valore, (int, float)) and not isinstance(valore, bool):
                    temperatura = float(valore)
                    allarme = temperatura < minimo or temperatura > massimo
                    record = {
                        **base,
                        "temp": temperatura,
                        "allarme": allarme,
                        "soglie": {"min": minimo, "max": massimo},
                        "origine_formula_excel": bool(is_formula),
                    }
                    formule += int(is_formula)
                    anomalie += int(allarme)
                else:
                    raise HTTPException(
                        422,
                        f"{filename}, foglio {nome}, {data.isoformat()}: valore non riconosciuto {valore!r}",
                    )
                temperature[str(mese)][str(giorno)] = record
                totale += 1

        docs.append({
            "id": str(uuid.uuid4()),
            "anno": anno,
            campo_numero: numero,
            campo_nome: nome,
            "azienda": "Ceraldi Group S.R.L.",
            "indirizzo": "Piazza Carità 14, 80134 Napoli (NA)",
            "piva": "04523831214",
            "temperature": temperature,
            "temp_min": minimo,
            "temp_max": massimo,
            "created_at": datetime.now(timezone.utc).isoformat(),
            "updated_at": datetime.now(timezone.utc).isoformat(),
            "importazione_excel": {
                "file": filename,
                "sha256": digest,
                "importato_il": datetime.now(timezone.utc).isoformat(),
                "formule_volatili": formule,
            },
        })

    return {
        "anno": anno,
        "tipo": tipo,
        "file": filename,
        "sha256": digest,
        "nomi": wf.sheetnames,
        "nomi_generici": generici,
        "documenti": docs,
        "totale": totale,
        "formule": formule,
        "chiusi": chiusi,
        "anomalie": anomalie,
        "mancanti": mancanti,
    }


def leggi_sanificazioni(raw: bytes, filename: str, oggi: date | None = None) -> dict:
    oggi = oggi or datetime.now(FUSO).date()
    wf, wv = _workbooks(raw)
    anno = _anno(wf)
    if anno != oggi.year:
        raise HTTPException(422, f"{filename}: anno {anno}, atteso {oggi.year}")
    digest = _sha256(raw)
    docs = []
    totale = chiusi = formule_orario = 0

    for ws_index, sf in enumerate(wf.worksheets):
        sv = wv.worksheets[ws_index]
        mese = MESI.get(sf.title.strip().upper())
        if not mese:
            continue
        registrazioni: dict[str, dict[str, str]] = {}
        firme: dict[str, dict[str, dict[str, Any]]] = {}
        originali: dict[str, dict[str, str]] = {}
        colonne = []
        for col in range(2, 8):
            intestazione = str(sf.cell(6, col).value or "").strip()
            area = SANIFICAZIONE_ALIASES.get(intestazione.upper(), intestazione)
            if area:
                colonne.append((col, area))
                registrazioni.setdefault(area, {})
                firme.setdefault(area, {})
                originali.setdefault(area, {})

        for row in range(9, min(sf.max_row, 39) + 1):
            giorno = sf.cell(row, 1).value
            if not isinstance(giorno, int) or not 1 <= giorno <= calendar.monthrange(anno, mese)[1]:
                continue
            data = date(anno, mese, giorno)
            if data > oggi:
                continue
            operatore = str(sf.cell(row, 9).value or "").strip()
            ora_formula = sf.cell(row, 8).value
            ora_salvata = sv.cell(row, 8).value if isinstance(ora_formula, str) and ora_formula.startswith("=") else ora_formula
            formule_orario += int(isinstance(ora_formula, str) and ora_formula.startswith("="))
            for col, area in colonne:
                valore = sf.cell(row, col).value
                if valore is None or str(valore).strip() == "":
                    continue
                originale = str(valore).strip()
                normalizzato = originale.upper()
                if normalizzato in {"CHIUSO", "CHIUSI", "C"}:
                    registrato = "N/D"
                    chiusi += 1
                elif normalizzato in {"X", "XX"}:
                    registrato = "X"
                    totale += 1
                else:
                    raise HTTPException(
                        422,
                        f"{filename}, foglio {sf.title}, giorno {giorno}: valore {originale!r} non riconosciuto",
                    )
                registrazioni[area][str(giorno)] = registrato
                originali[area][str(giorno)] = originale
                if registrato == "X":
                    firme[area][str(giorno)] = {
                        "valore": registrato,
                        "valore_originale": originale,
                        "operatore": operatore,
                        "firma_verificata": False,
                        "firma_via": "import_excel",
                        "ora_dichiarata": ora_salvata,
                        "timestamp": _timestamp_giorno(data),
                    }

        docs.append({
            "id": str(uuid.uuid4()),
            "mese": mese,
            "anno": anno,
            "azienda": "Ceraldi Group S.R.L.",
            "indirizzo": "Piazza Carità 14 Napoli",
            "area": "Sala e Servizi",
            "registrazioni": registrazioni,
            "firme": firme,
            "valori_originali_excel": originali,
            "operatore_responsabile": "",
            "created_at": datetime.now(timezone.utc).isoformat(),
            "updated_at": datetime.now(timezone.utc).isoformat(),
            "importazione_excel": {"file": filename, "sha256": digest},
        })

    if not docs:
        raise HTTPException(422, f"{filename}: nessun foglio mensile riconosciuto")
    return {
        "anno": anno,
        "file": filename,
        "sha256": digest,
        "documenti": docs,
        "totale": totale,
        "chiusi": chiusi,
        "formule_orario": formule_orario,
    }


def prepara_importazione(neg_raw: bytes, neg_name: str, pos_raw: bytes, pos_name: str,
                         san_raw: bytes, san_name: str, oggi: date | None = None) -> dict:
    negativo = leggi_temperature(neg_raw, neg_name, "congelatore", oggi)
    positivo = leggi_temperature(pos_raw, pos_name, "frigo", oggi)
    sanificazione = leggi_sanificazioni(san_raw, san_name, oggi)
    anni = {negativo["anno"], positivo["anno"], sanificazione["anno"]}
    if len(anni) != 1:
        raise HTTPException(422, "I tre file non appartengono allo stesso anno")
    avvisi = []
    if negativo["formule"] or positivo["formule"]:
        avvisi.append(
            "Le temperature provengono in gran parte da formule Excel volatili RANDBETWEEN; "
            "vengono importati i valori salvati, senza ricalcolarli."
        )
    if negativo["anomalie"] or positivo["anomalie"]:
        avvisi.append(
            f"Il file contiene {negativo['anomalie'] + positivo['anomalie']} valori fuori dal range HACCP."
        )
    if negativo["nomi_generici"] or positivo["nomi_generici"]:
        avvisi.append("Sono presenti nomi foglio generici: " + ", ".join(
            negativo["nomi_generici"] + positivo["nomi_generici"]
        ))
    return {
        "anno": anni.pop(),
        "negative": negativo,
        "positive": positivo,
        "sanificazione": sanificazione,
        "avvisi": avvisi,
    }


def riepilogo(preparata: dict) -> dict:
    return {
        "anno": preparata["anno"],
        "negative": {k: preparata["negative"][k] for k in (
            "file", "sha256", "nomi", "nomi_generici", "totale", "formule", "chiusi", "anomalie", "mancanti"
        )},
        "positive": {k: preparata["positive"][k] for k in (
            "file", "sha256", "nomi", "nomi_generici", "totale", "formule", "chiusi", "anomalie", "mancanti"
        )},
        "sanificazione": {k: preparata["sanificazione"][k] for k in (
            "file", "sha256", "totale", "chiusi", "formule_orario"
        )},
        "avvisi": preparata["avvisi"],
    }


async def sostituisci_archivio(db, preparata: dict, attore: dict | None = None) -> dict:
    """Crea un backup completo e sostituisce i dati; in errore ripristina il backup."""
    anno = preparata["anno"]
    import_id = str(uuid.uuid4())
    timestamp = datetime.now(timezone.utc).isoformat()
    backup = {
        "id": import_id,
        "anno": anno,
        "created_at": timestamp,
        "created_by": (attore or {}).get("nome") or (attore or {}).get("id") or "",
        "fonti": {
            "negative": preparata["negative"]["sha256"],
            "positive": preparata["positive"]["sha256"],
            "sanificazione": preparata["sanificazione"]["sha256"],
        },
        "temperature_positive": await db.temperature_positive.find({"anno": anno}, {"_id": 0}).to_list(None),
        "temperature_negative": await db.temperature_negative.find({"anno": anno}, {"_id": 0}).to_list(None),
        "sanificazione_schede": await db.sanificazione_schede.find({"anno": anno}, {"_id": 0}).to_list(None),
        "attrezzature_config": await db.attrezzature_config.find(
            {"tipo": {"$in": ["frigo", "congelatore"]}}, {"_id": 0}
        ).to_list(None),
        "stato": "backup_creato",
    }
    await db.haccp_import_backups.insert_one(dict(backup))

    async def _ripristina():
        for coll_name, key in (
            ("temperature_positive", "temperature_positive"),
            ("temperature_negative", "temperature_negative"),
            ("sanificazione_schede", "sanificazione_schede"),
        ):
            coll = getattr(db, coll_name)
            await coll.delete_many({"anno": anno})
            if backup[key]:
                await coll.insert_many(backup[key])
        await db.attrezzature_config.delete_many({"tipo": {"$in": ["frigo", "congelatore"]}})
        if backup["attrezzature_config"]:
            await db.attrezzature_config.insert_many(backup["attrezzature_config"])

    try:
        await db.temperature_positive.delete_many({"anno": anno})
        await db.temperature_negative.delete_many({"anno": anno})
        await db.sanificazione_schede.delete_many({"anno": anno})
        await db.temperature_positive.insert_many(preparata["positive"]["documenti"])
        await db.temperature_negative.insert_many(preparata["negative"]["documenti"])
        await db.sanificazione_schede.insert_many(preparata["sanificazione"]["documenti"])

        for tipo, fonte in (("frigo", preparata["positive"]), ("congelatore", preparata["negative"])):
            numeri = list(range(1, len(fonte["nomi"]) + 1))
            await db.attrezzature_config.update_many(
                {"tipo": tipo, "numero": {"$nin": numeri}},
                {"$set": {"attivo": False, "disattivato_da_import": import_id, "updated_at": timestamp}},
            )
            for numero, nome in enumerate(fonte["nomi"], 1):
                await db.attrezzature_config.update_one(
                    {"tipo": tipo, "numero": numero},
                    {"$set": {"tipo": tipo, "numero": numero, "nome": nome, "attivo": True,
                              "updated_at": timestamp, "import_id": import_id}},
                    upsert=True,
                )
        await db.haccp_import_backups.update_one(
            {"id": import_id}, {"$set": {"stato": "applicato", "applied_at": timestamp}}
        )
    except Exception:
        await _ripristina()
        await db.haccp_import_backups.update_one(
            {"id": import_id}, {"$set": {"stato": "ripristinato_dopo_errore"}}
        )
        raise

    return {"import_id": import_id, "backup_id": import_id, **riepilogo(preparata)}

