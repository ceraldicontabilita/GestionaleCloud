"""Listino di un fornitore da file (Excel o CSV) → catalogo di Lotti.

Un listino e' il prezzo che il fornitore **dichiara** oggi al cliente (per
Barone: il catalogo riservato a Ceraldi Group scaricato dal suo sito). Non e'
un acquisto: nel confronto prezzi compare con la sua data e la scritta
«listino», accanto ai prezzi pagati in fattura XML, mai mescolato con loro.

Il file si legge per **intestazione di colonna** (Codice, Descrizione, EAN,
IVA, Unita' di misura, Prezzo, Categoria...), non per posizione: un fornitore
diverso con le colonne in un altro ordine si carica lo stesso. Una riga senza
codice, descrizione o prezzo positivo si scarta e si conta, non si indovina.

I prodotti vanno nella collezione dei cataloghi (``catalogo_forno_prodotti``,
chiave fornitore + codice articolo), la fonte nel registro dei cataloghi
(``fonti_catalogo_esterne``, ``tipo="listino"``): la stessa scheda del catalogo
in Prodotti, nessun sistema parallelo. Un articolo che sparisce dal listino
nuovo non si cancella: ``nel_listino=False`` con la data.

Importi ``Decimal``, salvati come stringa (``"8.52"``).
"""
from __future__ import annotations

from app.lotti.bulk_compat import bulk_write_compat

import csv
import hashlib
import io
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from typing import Any, Dict, Iterable, List, Optional, Tuple

COLLEZIONE_PRODOTTI = "catalogo_forno_prodotti"
COLLEZIONE_FONTI = "fonti_catalogo_esterne"

# intestazione normalizzata → campo
_COLONNE = {
    "codice": "codice", "cod": "codice", "codice articolo": "codice", "cod articolo": "codice",
    "articolo": "codice", "sku": "codice",
    "descrizione": "descrizione", "descrizione articolo": "descrizione", "prodotto": "descrizione",
    "nome": "descrizione", "denominazione": "descrizione",
    "ean": "ean", "barcode": "ean", "codice a barre": "ean", "gtin": "ean",
    "iva": "iva", "iva %": "iva", "aliquota iva": "iva", "aliquota": "iva", "cod iva": "iva",
    "unita di misura": "unita", "um": "unita", "u m": "unita", "unita": "unita", "unita misura": "unita",
    "prezzo": "prezzo", "prezzo €": "prezzo", "prezzo eur": "prezzo", "prezzo netto": "prezzo",
    "prezzo listino": "prezzo", "listino": "prezzo", "prezzo unitario": "prezzo",
    "conf ordine": "confezione", "confezione": "confezione", "conf": "confezione",
    "categoria": "categoria", "reparto": "categoria", "categoria merceologica": "categoria",
    "sottocategoria": "sottocategoria", "sotto categoria": "sottocategoria", "famiglia": "sottocategoria",
    "link scheda": "link", "link": "link", "url": "link", "scheda": "link",
    "offerta volantino fino al": "offerta_fino", "offerta fino al": "offerta_fino", "fine offerta": "offerta_fino",
}
_OBBLIGATORIE = ("codice", "descrizione", "prezzo")

_RX_UNITA = re.compile(r"^\s*(\d+(?:[.,]\d+)?)?\s*([A-Za-z.]+)?\s*$")


def _norm_intestazione(testo: Any) -> str:
    t = unicodedata.normalize("NFKD", str(testo or ""))
    t = "".join(c for c in t if not unicodedata.combining(c)).lower()
    t = t.replace("€", " € ")
    t = re.sub(r"[^a-z0-9%€ ]+", " ", t)
    return re.sub(r"\s+", " ", t).strip()


def decimale(valore: Any) -> Optional[Decimal]:
    """«8,52», «8.52», «1.234,56», 8.52 → Decimal; vuoto o illeggibile → None."""
    if valore is None or valore == "":
        return None
    if isinstance(valore, (int, float, Decimal)) and not isinstance(valore, bool):
        try:
            return Decimal(str(valore))
        except InvalidOperation:
            return None
    testo = str(valore).strip().replace("€", "").replace(" ", "")
    if not testo:
        return None
    if "," in testo:
        testo = testo.replace(".", "").replace(",", ".")
    try:
        return Decimal(testo)
    except InvalidOperation:
        return None


def leggi_unita(testo: Any) -> Tuple[int, str]:
    """«12 PZ» → (12, "PZ"); «PZ» → (1, "PZ"); «KG» → (1, "KG"); «6 CF» → (6, "CF")."""
    m = _RX_UNITA.match(str(testo or "").strip().upper())
    if not m:
        return 1, str(testo or "").strip().upper()
    numero, sigla = m.groups()
    n = 1
    if numero:
        valore = decimale(numero)
        if valore and valore == valore.to_integral_value() and valore > 0:
            n = int(valore)
    return n, (sigla or "").rstrip(".")


def slug(nome: str) -> str:
    s = re.sub(r"[^a-z0-9]+", "", unicodedata.normalize("NFKD", nome or "").encode("ascii", "ignore").decode().lower())
    return s or "fornitore"


def partita_iva(testo: Any) -> str:
    p = re.sub(r"[^0-9A-Z]", "", str(testo or "").upper())
    return re.sub(r"^IT(?=\d{11}$)", "", p)


@dataclass
class RigaListino:
    codice: str
    descrizione: str
    prezzo: Decimal
    unita: str = ""               # «12 PZ» come scritto dal fornitore
    pezzi_vendita: int = 1        # 12
    sigla_unita: str = ""         # «PZ», «KG», «CF»
    ean: str = ""
    iva: Optional[int] = None
    confezione: str = ""
    categoria: str = ""
    sottocategoria: str = ""
    link: str = ""
    offerta_fino: str = ""        # «2026-12-31»: prezzo in offerta volantino


@dataclass
class EsitoLettura:
    righe: List[RigaListino] = field(default_factory=list)
    scartate: List[Dict[str, Any]] = field(default_factory=list)
    colonne: Dict[str, str] = field(default_factory=dict)


def _mappa_colonne(intestazione: Iterable[Any]) -> Dict[int, str]:
    mappa: Dict[int, str] = {}
    for i, cella in enumerate(intestazione):
        campo = _COLONNE.get(_norm_intestazione(cella))
        if campo and campo not in mappa.values():
            mappa[i] = campo
    return mappa


def _testo(valore: Any) -> str:
    if valore is None:
        return ""
    if isinstance(valore, float) and valore.is_integer():
        valore = int(valore)
    return re.sub(r"\s+", " ", str(valore)).strip()


def leggi_righe(tabella: Iterable[Iterable[Any]]) -> EsitoLettura:
    """Tabella grezza (prima riga utile = intestazione) → righe del listino."""
    esito = EsitoLettura()
    mappa: Dict[int, str] = {}
    for n_riga, riga in enumerate(tabella, start=1):
        celle = list(riga)
        if not mappa:
            candidata = _mappa_colonne(celle)
            if all(c in candidata.values() for c in _OBBLIGATORIE):
                mappa = candidata
                esito.colonne = {campo: _testo(celle[i]) for i, campo in mappa.items()}
            continue
        valori = {campo: celle[i] if i < len(celle) else None for i, campo in mappa.items()}
        if not any(_testo(v) for v in valori.values()):
            continue
        codice = _testo(valori.get("codice"))
        descrizione = _testo(valori.get("descrizione"))
        prezzo = decimale(valori.get("prezzo"))
        motivo = ""
        if not codice:
            motivo = "codice mancante"
        elif not descrizione:
            motivo = "descrizione mancante"
        elif prezzo is None or prezzo <= 0:
            motivo = "prezzo mancante o zero"
        if motivo:
            esito.scartate.append({"riga": n_riga, "codice": codice, "descrizione": descrizione, "motivo": motivo})
            continue
        unita = _testo(valori.get("unita")).upper()
        pezzi, sigla = leggi_unita(unita)
        iva = decimale(valori.get("iva"))
        ean = re.sub(r"\D", "", _testo(valori.get("ean")))
        esito.righe.append(RigaListino(
            codice=codice,
            descrizione=descrizione,
            prezzo=prezzo,
            unita=unita,
            pezzi_vendita=pezzi,
            sigla_unita=sigla,
            ean=ean if 8 <= len(ean) <= 14 else "",
            iva=int(iva) if iva is not None and iva == iva.to_integral_value() else None,
            confezione=_testo(valori.get("confezione")).upper(),
            categoria=_testo(valori.get("categoria")),
            sottocategoria=_testo(valori.get("sottocategoria")),
            link=_testo(valori.get("link")),
            offerta_fino=_data_iso(valori.get("offerta_fino")),
        ))
    if not mappa:
        raise ValueError(
            "Intestazione non trovata: servono almeno le colonne Codice, Descrizione e Prezzo"
        )
    # stesso codice due volte: vale la prima, la seconda si segnala
    visti: Dict[str, int] = {}
    uniche: List[RigaListino] = []
    for r in esito.righe:
        if r.codice in visti:
            esito.scartate.append({"codice": r.codice, "descrizione": r.descrizione, "motivo": "codice ripetuto"})
            continue
        visti[r.codice] = 1
        uniche.append(r)
    esito.righe = uniche
    return esito


def tabella_da_file(contenuto: bytes, nome_file: str) -> List[List[Any]]:
    """Excel (.xlsx) o CSV (; o ,) → righe di celle. Per un Excel si legge il
    foglio che contiene l'intestazione (il primo, se nessuno la contiene)."""
    nome = (nome_file or "").lower()
    if nome.endswith(".csv") or nome.endswith(".txt"):
        testo = contenuto.decode("utf-8-sig", errors="replace")
        dialetto = csv.Sniffer().sniff(testo[:4096], delimiters=";,\t")
        return [list(r) for r in csv.reader(io.StringIO(testo), dialetto)]
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(contenuto), read_only=True, data_only=True)
    try:
        fogli = [[list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets]
    finally:
        wb.close()
    for righe in fogli:
        for riga in righe[:20]:
            candidata = _mappa_colonne(riga)
            if all(c in candidata.values() for c in _OBBLIGATORIE):
                return righe
    return fogli[0] if fogli else []


def _data_iso(valore: Any) -> str:
    if isinstance(valore, datetime):
        return valore.date().isoformat()
    testo = _testo(valore)[:10]
    for formato in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(testo, formato).date().isoformat()
        except ValueError:
            continue
    return ""


def sha256(contenuto: bytes) -> str:
    return hashlib.sha256(contenuto).hexdigest()


def documento_prodotto(r: RigaListino, *, fornitore_key: str, fornitore_nome: str,
                       data_listino: str, file_sha256: str, adesso: str) -> Dict[str, Any]:
    doc_id = f"listino:{fornitore_key}:{r.codice}"
    return {
        "_id": doc_id,
        "id": doc_id,
        "fornitore": fornitore_key,
        "fornitore_nome": fornitore_nome,
        "codice_articolo": r.codice,
        "nome": r.descrizione,
        "nome_completo": r.descrizione,
        "categoria": r.categoria,
        "sottocategoria": r.sottocategoria,
        "ean": r.ean,
        "aliquota_iva": r.iva,
        "unita_vendita": r.unita,
        "pezzi_vendita": r.pezzi_vendita,
        "sigla_unita": r.sigla_unita,
        "confezione_ordine": r.confezione,
        "prezzo_listino": f"{r.prezzo:f}",
        "prezzo_listino_valuta": "EUR",
        "prezzo_listino_iva_esclusa": True,
        "prezzo_listino_data": data_listino,
        "listino_sha256": file_sha256,
        "link_prodotto": r.link,
        "offerta_fino": r.offerta_fino,
        "nel_listino": True,
        "fonte_catalogo": "listino",
        "data_aggiornamento": adesso,
    }


_CAMPI_CONFRONTO = ("nome", "categoria", "sottocategoria", "ean", "aliquota_iva", "unita_vendita",
                    "prezzo_listino", "link_prodotto", "nel_listino", "confezione_ordine", "offerta_fino")


async def importa(db, esito: EsitoLettura, *, fornitore_nome: str, fornitore_key: str = "",
                  piva: str = "", url: str = "", data_listino: str = "", file_sha256: str = "",
                  nome_file: str = "") -> Dict[str, Any]:
    """Scrive il listino: upsert per (fornitore, codice), usciti marcati, fonte
    registrata. Il secondo import dello stesso file da' ``nuovi=0`` e
    ``aggiornati=0``."""
    from pymongo import UpdateOne

    nome = (fornitore_nome or "").strip()
    if not nome:
        raise ValueError("Il nome del fornitore e' obbligatorio")
    key = slug(fornitore_key or nome)
    adesso = datetime.now(timezone.utc).isoformat()
    data_listino = data_listino or adesso[:10]
    prodotti = getattr(db, COLLEZIONE_PRODOTTI)
    presenti = {
        str(d.get("codice_articolo")): d
        for d in await prodotti.find({"fornitore": key}, {"_id": 0}).to_list(None)
    }
    operazioni = []
    nuovi = aggiornati = invariati = 0
    codici = set()
    for r in esito.righe:
        codici.add(r.codice)
        doc = documento_prodotto(r, fornitore_key=key, fornitore_nome=nome, data_listino=data_listino,
                                 file_sha256=file_sha256, adesso=adesso)
        prima = presenti.get(r.codice)
        if prima is None:
            nuovi += 1
        elif all(prima.get(c) == doc.get(c) for c in _CAMPI_CONFRONTO):
            invariati += 1
            continue
        else:
            aggiornati += 1
            if prima.get("prezzo_listino") and prima.get("prezzo_listino") != doc["prezzo_listino"]:
                doc["prezzo_listino_precedente"] = prima.get("prezzo_listino")
                doc["prezzo_listino_precedente_data"] = prima.get("prezzo_listino_data")
        doc.pop("_id")
        operazioni.append(UpdateOne({"fornitore": key, "codice_articolo": r.codice},
                                    {"$set": doc, "$setOnInsert": {"_id": f"listino:{key}:{r.codice}"}},
                                    upsert=True))
    usciti = 0
    for codice, prima in presenti.items():
        if codice not in codici and prima.get("fonte_catalogo") == "listino" and prima.get("nel_listino") is not False:
            usciti += 1
            operazioni.append(UpdateOne({"fornitore": key, "codice_articolo": codice},
                                        {"$set": {"nel_listino": False, "uscito_dal_listino_il": data_listino}}))
    if operazioni:
        await bulk_write_compat(prodotti, operazioni, ordered=False)

    fonti = getattr(db, COLLEZIONE_FONTI)
    fonte = await fonti.find_one({"fornitore_key": key}, {"_id": 0})
    dati_fonte = {
        "nome": nome,
        "fornitore_key": key,
        "tipo": "listino",
        "partita_iva": partita_iva(piva) or (fonte or {}).get("partita_iva", ""),
        "stato": "attivo",
        "prodotti_trovati": len(esito.righe),
        "ultima_sincronizzazione": adesso,
        "listino_data": data_listino,
        "listino_sha256": file_sha256,
        "listino_file": nome_file,
        "ultimo_errore": None,
    }
    if url:
        dati_fonte["url"] = url
    if fonte is None:
        import uuid

        dati_fonte.update({"id": str(uuid.uuid4()), "creato_il": adesso, "url": url})
        await fonti.insert_one(dati_fonte)
    else:
        await fonti.update_one({"fornitore_key": key}, {"$set": dati_fonte})
    return {
        "ok": True,
        "fornitore": nome,
        "fornitore_key": key,
        "righe_lette": len(esito.righe),
        "nuovi": nuovi,
        "aggiornati": aggiornati,
        "invariati": invariati,
        "usciti_dal_listino": usciti,
        "scartate": len(esito.scartate),
        "esempi_scartate": esito.scartate[:10],
        "data_listino": data_listino,
    }
