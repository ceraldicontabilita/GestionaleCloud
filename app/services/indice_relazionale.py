"""Indice relazionale (MINI-06): una vista di sola lettura su ``entity_relations``
con esportazione in CSV, JSON e XLSX.

Non e' un secondo registro: legge le relazioni che esistono (fra cui i
documenti di DRV-03, target ``documento``), le appiattisce in righe leggibili e
le ordina per ``relation_key``. L'esportazione e' **riproducibile**: stessi dati
in archivio, stessi byte (nessuna data di generazione dentro il file, righe
ordinate, campi in ordine fisso, ZIP dell'XLSX con data fissa). Importi in
``Decimal`` dai centesimi; date ``gg/mm/aaaa`` in CSV e XLSX, ISO nel JSON.
"""
from __future__ import annotations

import csv
import io
import json
import zipfile
from collections import Counter
from datetime import datetime
from decimal import Decimal
from typing import Any, Dict, List, Optional, Sequence, Tuple

from app.db_collections import COLL_ENTITY_RELATIONS

FORMATI = ("csv", "json", "xlsx")
LIMITE_LETTURA = 100_000
ETICHETTA_STATO = {"confirmed": "CONFERMATA", "pending": "DA_VERIFICARE", "revoked": "REVOCATA"}

# (chiave, intestazione)
COLONNE: Tuple[Tuple[str, str], ...] = (
    ("relation_key", "Chiave"),
    ("origine_tipo", "Origine tipo"),
    ("origine_id", "Origine id"),
    ("relazione", "Relazione"),
    ("destinazione_tipo", "Destinazione tipo"),
    ("destinazione_id", "Destinazione id"),
    ("stato", "Stato"),
    ("regola", "Regola"),
    ("importo", "Importo EUR"),
    ("prove", "Prove"),
    ("creata", "Creata"),
    ("aggiornata", "Aggiornata"),
)

PROIEZIONE = {"_id": 0, "relation_key": 1, "source": 1, "target": 1, "relation_type": 1, "status": 1,
              "rule": 1, "evidence": 1, "amount_cents": 1, "created_at": 1, "updated_at": 1, "provenance": 1}


def _giorno(iso: Any) -> str:
    """ISO-8601 -> gg/mm/aaaa (vuoto se assente o illeggibile: mai una data inventata)."""
    testo = str(iso or "").strip()
    if not testo:
        return ""
    try:
        return datetime.fromisoformat(testo.replace("Z", "+00:00")).strftime("%d/%m/%Y")
    except ValueError:
        return ""


def _euro(cents: Any) -> Optional[Decimal]:
    if cents is None or isinstance(cents, bool):
        return None
    try:
        return (Decimal(int(cents)) / Decimal(100)).quantize(Decimal("0.01"))
    except (TypeError, ValueError):
        return None


def riga_indice(doc: Dict[str, Any]) -> Dict[str, Any]:
    """Una relazione appiattita. I campi sono sempre gli stessi, nello stesso ordine."""
    origine, destinazione = doc.get("source") or {}, doc.get("target") or {}
    prove = sorted(f"{e.get('type')}={e.get('value')}" for e in (doc.get("evidence") or []) if isinstance(e, dict))
    provenienza = doc.get("provenance") or {}
    return {
        "relation_key": doc.get("relation_key") or "",
        "origine_tipo": origine.get("type") or "",
        "origine_id": str(origine.get("id") or ""),
        "relazione": doc.get("relation_type") or "",
        "destinazione_tipo": destinazione.get("type") or "",
        "destinazione_id": str(destinazione.get("id") or ""),
        "stato": ETICHETTA_STATO.get(str(doc.get("status") or "").lower(), str(doc.get("status") or "").upper()),
        "regola": doc.get("rule") or "",
        "importo": _euro(doc.get("amount_cents")),
        "prove": "; ".join(prove),
        "creata": doc.get("created_at") or "",
        "aggiornata": doc.get("updated_at") or "",
        "motivo": provenienza.get("motivo") if isinstance(provenienza, dict) else None,
    }


def costruisci_filtro(origine_tipo: Optional[str] = None, destinazione_tipo: Optional[str] = None,
                      relazione: Optional[str] = None, stato: Optional[str] = None,
                      entita_id: Optional[str] = None) -> Dict[str, Any]:
    filtro: Dict[str, Any] = {}
    if origine_tipo:
        filtro["source.type"] = origine_tipo
    if destinazione_tipo:
        filtro["target.type"] = destinazione_tipo
    if relazione:
        filtro["relation_type"] = relazione
    if stato:
        # accetta sia il vocabolario dell'archivio sia le etichette dell'indice
        inverso = {v: k for k, v in ETICHETTA_STATO.items()}
        filtro["status"] = inverso.get(stato.upper(), stato.lower())
    if entita_id:
        filtro["$or"] = [{"source.id": entita_id}, {"target.id": entita_id}]
    return filtro


async def leggi_righe(db, filtro: Optional[Dict[str, Any]] = None, *, limite: int = LIMITE_LETTURA) -> Dict[str, Any]:
    """Tutte le relazioni che passano il filtro, senza payload, ordinate per chiave."""
    docs = await db[COLL_ENTITY_RELATIONS].find(filtro or {}, PROIEZIONE).limit(limite + 1).to_list(length=limite + 1)
    troncata = len(docs) > limite
    righe = sorted((riga_indice(d) for d in docs[:limite]), key=lambda r: r["relation_key"])
    return {"righe": righe, "troncata": troncata}


def riepilogo(righe: Sequence[Dict[str, Any]]) -> Dict[str, Any]:
    per_stato = Counter(r["stato"] for r in righe)
    per_coppia = Counter(f"{r['origine_tipo']} > {r['destinazione_tipo']}" for r in righe)
    return {"totale": len(righe), "per_stato": dict(sorted(per_stato.items())),
            "per_coppia": dict(sorted(per_coppia.items()))}


# ── esportazione ─────────────────────────────────────────────────────────────

def _cella_testo(valore: Any) -> str:
    testo = "" if valore is None else str(valore)
    # un campo che inizia con = + - @ sarebbe letto da Excel come formula
    return "'" + testo if testo[:1] in ("=", "+", "-", "@") else testo


def _valori_tabella(righe: Sequence[Dict[str, Any]]) -> List[List[Any]]:
    out: List[List[Any]] = []
    for r in righe:
        out.append([
            _cella_testo(r["relation_key"]), _cella_testo(r["origine_tipo"]), _cella_testo(r["origine_id"]),
            _cella_testo(r["relazione"]), _cella_testo(r["destinazione_tipo"]), _cella_testo(r["destinazione_id"]),
            r["stato"], _cella_testo(r["regola"]), r["importo"], _cella_testo(r["prove"]),
            _giorno(r["creata"]), _giorno(r["aggiornata"]),
        ])
    return out


def esporta_csv(righe: Sequence[Dict[str, Any]]) -> bytes:
    """UTF-8 con BOM (Excel), separatore `;`, importi con la virgola, fine riga `\\n`."""
    buf = io.StringIO()
    scrittore = csv.writer(buf, delimiter=";", lineterminator="\n")
    scrittore.writerow([intestazione for _, intestazione in COLONNE])
    for valori in _valori_tabella(righe):
        scrittore.writerow(["" if v is None else (f"{v:.2f}".replace(".", ",") if isinstance(v, Decimal) else v)
                            for v in valori])
    return b"\xef\xbb\xbf" + buf.getvalue().encode("utf-8")


def esporta_json(righe: Sequence[Dict[str, Any]]) -> bytes:
    """Date ISO com'e' nell'archivio, importo come stringa Decimal (mai float)."""
    voci = []
    for r in righe:
        voce = {chiave: r[chiave] for chiave, _ in COLONNE}
        voce["importo"] = None if r["importo"] is None else f"{r['importo']:.2f}"
        voce["motivo"] = r.get("motivo")
        voci.append(voce)
    return (json.dumps({"colonne": [c for c, _ in COLONNE], "relazioni": voci}, ensure_ascii=False,
                       sort_keys=True, indent=1) + "\n").encode("utf-8")


def esporta_xlsx(righe: Sequence[Dict[str, Any]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Indice relazionale"
    ws.append([intestazione for _, intestazione in COLONNE])
    for cella in ws[1]:
        cella.font = Font(bold=True)
    for valori in _valori_tabella(righe):
        ws.append(valori)
    colonna_importo = [c for c, _ in COLONNE].index("importo") + 1
    for riga in ws.iter_rows(min_row=2, min_col=colonna_importo, max_col=colonna_importo):
        riga[0].number_format = "#,##0.00"
    ws.freeze_panes = "A2"
    wb.properties.creator = "GestionaleCloud"
    wb.properties.created = wb.properties.modified = datetime(2000, 1, 1)
    grezzo = io.BytesIO()
    wb.save(grezzo)
    # openpyxl scrive la data corrente nelle voci dello ZIP: si riscrivono a data fissa
    fisso = io.BytesIO()
    with zipfile.ZipFile(grezzo) as sorgente, zipfile.ZipFile(fisso, "w", zipfile.ZIP_DEFLATED) as dest:
        for info in sorted(sorgente.infolist(), key=lambda i: i.filename):
            voce = zipfile.ZipInfo(info.filename, date_time=(2000, 1, 1, 0, 0, 0))
            voce.compress_type = zipfile.ZIP_DEFLATED
            voce.external_attr = 0o600 << 16
            dest.writestr(voce, sorgente.read(info.filename))
    return fisso.getvalue()


TIPI_MIME = {
    "csv": "text/csv; charset=utf-8",
    "json": "application/json",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def esporta(righe: Sequence[Dict[str, Any]], formato: str) -> Tuple[bytes, str, str]:
    """(contenuto, tipo MIME, nome file). Il nome non porta la data: stesso archivio, stesso file."""
    formato = (formato or "").lower()
    if formato not in FORMATI:
        raise ValueError(f"formato non valido: {formato}")
    corpo = {"csv": esporta_csv, "json": esporta_json, "xlsx": esporta_xlsx}[formato](righe)
    return corpo, TIPI_MIME[formato], f"indice_relazionale.{formato}"


__all__ = ["FORMATI", "COLONNE", "costruisci_filtro", "leggi_righe", "riepilogo", "esporta", "riga_indice"]
