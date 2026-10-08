"""Scadenzario dei tributi in Excel: un versamento per riga, con F24, quietanza e ritardo.

Lo schema e' quello del foglio che manda il titolare (Data, Descrizione, rateazione /
regione / mese di riferimento, anno di riferimento, Codice tributo, Importo, F24 del
consulente, Quietanza) con in piu' scadenza, giorni di ritardo e prospetto contabile.

Non si calcola niente di nuovo: le righe sono quelle del registro F24 (``righe_modello``),
la quietanza e' quella che il registro aggancia al modello (``prove_modello``), la scadenza
e' la regola dello scadenzario dei tributi (festivi e proroga di agosto compresi) e il
prospetto e' quello agganciato in ``prospetti_contabili``. Un dato che manca resta vuoto:
niente data inventata, niente ritardo senza scadenza.

Tre fogli:

* **Scadenzario** — una riga per riga tributo di ogni modello F24 (importo = debito - credito,
  quindi un credito compensato e' negativo), dal versamento piu' recente;
* **Riepilogo** — per anno e codice, il versato di ogni mese di versamento;
* **Debito e credito** — le stesse righe con le due colonne separate.
"""
from __future__ import annotations

import io
import re
from datetime import date
from typing import Any, Dict, Iterable, List, Optional, Tuple

from app.services.originale_documento import url_originale
from app.services import f24_controllo_incrociato as registro_f24
from app.services.prospetti_contabili import COLL_PROSPETTI, CANONICA
from app.services.scadenzario_tributi import scadenza_da_regola

MESI_IT = ["gennaio", "febbraio", "marzo", "aprile", "maggio", "giugno", "luglio",
           "agosto", "settembre", "ottobre", "novembre", "dicembre"]

INTESTAZIONI = [
    "Data", "Descrizione", "rateazione, regione/provincia, mese rif.", "anno di riferimento",
    "Codice tributo", "Importo", "f24 consulente", "quietanza",
    "scadenza", "giorni di ritardo", "prospetto contabile",
]


def _data(testo: Optional[str]) -> Optional[date]:
    try:
        return date.fromisoformat(str(testo)[:10]) if testo else None
    except ValueError:
        return None


def _codice_cella(codice: str) -> Any:
    """Un codice numerico resta numero (come nel foglio del titolare), «DM10» resta testo."""
    return int(codice) if codice.isdigit() else codice


async def righe_scadenzario(db, anni: Optional[Iterable[int]] = None) -> List[Dict[str, Any]]:
    """Le righe del foglio. ``anni=None`` = tutti. Un modello conta una volta (copie uguali sono una)."""
    registro = await registro_f24.carica_registro(db)
    prospetti = await db[COLL_PROSPETTI].find({"stato": CANONICA, "f24_id": {"$ne": None}}, {"_id": 0}).to_list(5000)
    prospetto_per_f24 = {str(p["f24_id"]): p for p in prospetti}
    anni_set = set(anni) if anni is not None else None

    visti: set = set()
    righe: List[Dict[str, Any]] = []
    for f24 in registro["f24"]:
        chiave = (registro_f24.data_versamento_modello(f24), registro_f24.saldo_modello_cents(f24))
        if chiave != (None, None):
            if chiave in visti:
                continue
            visti.add(chiave)
        prove = registro_f24.prove_modello(f24, registro)
        agganciate = [q for q in prove["quietanze"] if q.get("quietanza_id")]
        quietanza = next((q for q in agganciate if q.get("data")), agganciate[0] if agganciate else None)
        data_v = _data((quietanza or {}).get("data")) or _data(prove["data_versamento"])
        prospetto = prospetto_per_f24.get(str(f24.get("id")))
        for r in registro_f24.righe_modello(f24):
            if not r["codice"]:
                continue
            anno_ref = r["anno"] if isinstance(r["anno"], int) else None
            # L'anno e' quello del versamento; senza data, quello di riferimento della riga.
            anno_filtro = data_v.year if data_v else anno_ref
            if anni_set is not None and anno_filtro not in anni_set:
                continue
            scadenza, _fonte = scadenza_da_regola(
                r["sezione"] or "", r["codice"], anno_ref,
                r["mese"] if isinstance(r["mese"], int) else None,
                data_v.isoformat() if data_v else None,
            )
            ritardo = max(0, (data_v - scadenza).days) if data_v and scadenza else None
            righe.append({
                "data": data_v, "descrizione": r["descrizione"], "mese": r["mese"], "anno": anno_ref,
                "periodo": r["periodo_riferimento"], "codice": r["codice"],
                "debito_cents": r["importo_debito_cents"], "credito_cents": r["importo_credito_cents"],
                "f24_id": f24.get("id"), "f24_data": _data(prove["data_versamento"]),
                "quietanza_id": (quietanza or {}).get("quietanza_id"),
                "quietanza_data": _data((quietanza or {}).get("data")),
                "scadenza": scadenza, "giorni_ritardo": ritardo,
                "prospetto": ({"periodo": f"{prospetto['mese']:02d}/{prospetto['anno']}", "esito": prospetto.get("esito"),
                               "documento_id": prospetto.get("documento_id")} if prospetto else None),
            })
    # Il versamento piu' recente per primo; senza data in fondo.
    righe.sort(key=lambda r: (r["data"] is None, -(r["data"].toordinal() if r["data"] else 0), r["codice"]))
    return righe


def _data_it(d: Optional[date]) -> str:
    return d.strftime("%d/%m/%Y") if d else ""


def costruisci_xlsx(righe: List[Dict[str, Any]], base_url: str = "") -> bytes:
    """Il file Excel. ``base_url`` e' l'indirizzo del gestionale (per i collegamenti)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    base = (base_url or "").rstrip("/")
    intestazione_font = Font(bold=True, color="FFFFFF")
    intestazione_fondo = PatternFill("solid", fgColor="0070C0")
    centro = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sinistra = Alignment(horizontal="left", vertical="top")
    EURO = '#,##0.00\\ "€";[Red]\\-#,##0.00\\ "€"'
    link_font = Font(color="0563C1", underline="single")

    def intesta(ws, titoli, larghezze):
        ws.append(titoli)
        for cella in ws[1]:
            cella.font, cella.fill, cella.alignment = intestazione_font, intestazione_fondo, centro
        for i, w in enumerate(larghezze, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(titoli))}1"

    wb = Workbook()
    ws = wb.active
    ws.title = "Scadenzario"
    intesta(ws, INTESTAZIONI, [14, 70, 22, 14, 14, 16, 22, 24, 14, 12, 28])
    for riga in righe:
        importo = (riga["debito_cents"] - riga["credito_cents"]) / 100
        ws.append([
            riga["data"], riga["descrizione"], riga["mese"] if riga["mese"] else (riga["periodo"] or None),
            riga["anno"], _codice_cella(riga["codice"]), importo,
            f"F24 {_data_it(riga['f24_data'])}".strip() if riga["f24_id"] else None,
            f"Quietanza {_data_it(riga['quietanza_data'])}".strip() if riga["quietanza_id"] else None,
            riga["scadenza"], riga["giorni_ritardo"],
            (f"{riga['prospetto']['periodo']} · {riga['prospetto']['esito'] or 'da agganciare'}" if riga["prospetto"] else None),
        ])
        n = ws.max_row
        ws.cell(n, 1).number_format = "dd/mm/yyyy"
        ws.cell(n, 9).number_format = "dd/mm/yyyy"
        ws.cell(n, 6).number_format = EURO
        for col in range(1, 12):
            ws.cell(n, col).alignment = sinistra
        if riga["f24_id"] and base:
            c = ws.cell(n, 7)
            c.hyperlink, c.font = f"{base}/fiscale/f24/{riga['f24_id']}", link_font
        if riga["quietanza_id"] and base:
            c = ws.cell(n, 8)
            c.hyperlink, c.font = f"{base}/fiscale/f24/{riga['quietanza_id']}", link_font
        if riga["prospetto"] and riga["prospetto"].get("documento_id") and base:
            c = ws.cell(n, 11)
            c.hyperlink, c.font = f"{base}{url_originale('documento', riga['prospetto']['documento_id'])}", link_font
        if riga["giorni_ritardo"]:
            ws.cell(n, 10).font = Font(bold=True, color="C00000")

    # Riepilogo: anno x codice x mese di versamento (le righe senza data non hanno un mese).
    ws2 = wb.create_sheet("Riepilogo")
    intesta(ws2, ["Anno", "Codice"] + [m for m in MESI_IT] + ["Totale"], [8, 12] + [13] * 12 + [15])
    somme: Dict[Tuple[int, Any], List[int]] = {}
    for riga in righe:
        if not riga["data"]:
            continue
        blocco = somme.setdefault((riga["data"].year, _codice_cella(riga["codice"])), [0] * 12)
        blocco[riga["data"].month - 1] += riga["debito_cents"] - riga["credito_cents"]
    for (anno, codice), mesi in sorted(somme.items(), key=lambda kv: (-kv[0][0], str(kv[0][1]))):
        ws2.append([anno, codice] + [m / 100 for m in mesi] + [sum(mesi) / 100])
        for col in range(3, 16):
            ws2.cell(ws2.max_row, col).number_format = EURO
    senza_data = [r for r in righe if not r["data"]]
    if senza_data:
        ws2.append([])
        ws2.append([f"{len(senza_data)} righe senza data di versamento: sono nel foglio Scadenzario, non hanno un mese."])

    # Debito e credito: le due colonne separate.
    ws3 = wb.create_sheet("Debito e credito")
    intesta(ws3, ["data", "tributo", "descrizione", "rateazione, regione/provincia, mese rif.",
                  "anno di riferimento", "importo a debito", "importo a credito"], [14, 12, 70, 22, 14, 18, 18])
    for riga in righe:
        ws3.append([riga["data"], _codice_cella(riga["codice"]), riga["descrizione"],
                    riga["mese"] if riga["mese"] else (riga["periodo"] or None), riga["anno"],
                    riga["debito_cents"] / 100, riga["credito_cents"] / 100])
        n = ws3.max_row
        ws3.cell(n, 1).number_format = "dd/mm/yyyy"
        ws3.cell(n, 6).number_format = EURO
        ws3.cell(n, 7).number_format = EURO
        for col in range(1, 8):
            ws3.cell(n, col).alignment = sinistra

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def nome_file(testo_anni: Optional[str]) -> str:
    pulito = re.sub(r"[^0-9a-z,-]", "", (testo_anni or "").lower()) or "anno"
    return f"scadenzario-tributi-{pulito}.xlsx"
