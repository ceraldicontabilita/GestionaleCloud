"""Protocollo personale e familiare (MINI-07): registro importato, ricerca AND, ponte informativo.

Il registro e' il foglio ``REGISTRO_PROTOCOLLO`` del file xlsx che il titolare
tiene su Drive (una riga per documento: numero ``AAAA/NNNNNN``, date, tipo,
oggetto letto dall'OCR, controparte, pratica, importo, nome file, SHA-256,
link). Qui diventa la collezione ``protocollo_personale``: una riga per numero.

Regole di questo modulo:

* **Il foglio si legge per intestazione**, mai per posizione: l'ordine delle
  colonne puo' cambiare, una colonna in piu' non rompe niente, una obbligatoria
  assente fa fallire l'import con il suo nome.
* **Il numero e' unico e immutabile.** Un numero gia' in archivio con un altro
  SHA-256 non si tocca: va in ``in_conflitto`` e lo decide il titolare. Le altre
  differenze (oggetto, note, pratica) aggiornano la riga.
* **Il protocollo non dimentica.** Niente si cancella: una riga tolta a mano
  diventa ``stato='rimosso'`` (con motivo e data) e resta cercabile su richiesta.
  Un file xlsx che non la contiene piu' non la rimuove, e un nuovo import non la
  riattiva.
* **Idempotente.** Lo stesso foglio due volte da' ``nuovi=0`` e ``aggiornati=0``.
  ``dry_run`` (predefinito) conta soltanto.
* **Ponte solo informativo** verso la contabilita': ``documenti_collegati``
  legge ``entity_relations`` e le impronte SHA-256 dei documenti gia' in archivio
  e li mostra. Non scrive mai una relazione, una scrittura o un pagamento: una
  riga personale/familiare resta fuori da bilanci, costi e Prima Nota.
* **Niente dati personali nei log**: solo contatori e numeri di riga.

Il testo intero di un PDF (``testo_ocr``) e' un payload e si legge per id; la
ricerca lavora su campi leggeri, fra cui ``testo_indice`` (i primi caratteri del
testo).
"""
from __future__ import annotations

import asyncio
import hashlib
import io
import logging
import re
import unicodedata
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

from app.constants.canale_documento import CANALE_DRIVE
from app.document_repository import metadata_projection

logger = logging.getLogger(__name__)

COLL = "protocollo_personale"
CHIAVE_STATO = "protocollo_personale_import"
FOGLIO = "REGISTRO_PROTOCOLLO"

STATO_ATTIVO = "attivo"
STATO_RIMOSSO = "rimosso"

AMBITO_FAMILIARE = "personale_familiare"
AMBITO_AZIENDALE = "aziendale"
AMBITO_DA_VERIFICARE = "da_verificare"

MAX_OGGETTO = 2000
MAX_TESTO_INDICE = 8000
MAX_TESTO_OCR = 400_000
MAX_ELENCO = 200
LIMITE_PAGINA = 200
LOTTO_SCRITTURA = 200

# intestazione normalizzata -> chiave interna (alias ammessi: mai la posizione)
INTESTAZIONI: Dict[str, str] = {
    "N_PROTOCOLLO": "numero", "NUMERO_PROTOCOLLO": "numero",
    "DATA_PROTOCOLLO": "data_protocollo",
    "TIPO_CORRISPONDENZA": "tipo_corrispondenza",
    "TIPO_DOCUMENTO": "tipo_documento",
    "DATA_DOCUMENTO": "data_documento",
    "ENTRATA_USCITA": "direzione",
    "OGGETTO": "oggetto",
    "MITTENTE_DESTINATARIO": "controparte",
    "PRATICA": "pratica",
    "IMPORTO": "importo",
    "NOTE": "note",
    "NOME_FILE": "nome_file",
    "IMPRONTA_SHA256": "sha256", "SHA256": "sha256", "SHA_256": "sha256",
    "LINK": "link",
}
OBBLIGATORIE = ("numero", "data_protocollo", "nome_file")

_NUMERO_RE = re.compile(r"^\s*(\d{4})\s*/\s*(\d{1,7})\s*$")
_SHA_RE = re.compile(r"^[0-9a-f]{64}$")
_DRIVE_ID_RE = re.compile(r"(?:/d/|[?&]id=|/folders/)([A-Za-z0-9_-]{10,})")
_NON_DATO = {"", "N/A", "N/D", "ND", "NA", "-", "--", "NONE", "NULL"}

_lavoro: Optional[asyncio.Task] = None
_lock = asyncio.Lock()


class RegistroNonValido(ValueError):
    """Il file non e' un registro leggibile (foglio o colonne obbligatorie assenti)."""

    def __init__(self, codice: str, messaggio: str, dettagli: Optional[Dict[str, Any]] = None):
        super().__init__(messaggio)
        self.codice = codice
        self.dettagli = dettagli or {}


# ── normalizzazione e campi ──────────────────────────────────────────────────

def _intestazione(valore: Any) -> str:
    return re.sub(r"[^A-Z0-9]+", "_", str(valore or "").strip().upper()).strip("_")


def _testo(valore: Any) -> Optional[str]:
    if valore is None:
        return None
    testo = " ".join(str(valore).split())
    return None if testo.upper() in _NON_DATO else testo


def numero_canonico(valore: Any) -> Optional[Tuple[str, int, int]]:
    """``AAAA/NNNNNN`` con lo zero-padding a 6 cifre; None se non e' un numero di protocollo."""
    trovato = _NUMERO_RE.match(str(valore or ""))
    if not trovato:
        return None
    anno, progressivo = int(trovato.group(1)), int(trovato.group(2))
    if anno < 1900 or progressivo < 1:
        return None
    return f"{anno:04d}/{progressivo:06d}", anno, progressivo


def data_iso(valore: Any) -> Optional[str]:
    """ISO ``AAAA-MM-GG`` da una data Excel, ISO o ``gg/mm/aaaa``; None se illeggibile (mai inventata)."""
    if isinstance(valore, datetime):
        return valore.date().isoformat()
    if isinstance(valore, date):
        return valore.isoformat()
    testo = str(valore or "").strip()
    if not testo or testo.upper() in _NON_DATO:
        return None
    for formato in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(testo[:10], formato).date().isoformat()
        except ValueError:
            continue
    return None


def importo_decimal(valore: Any) -> Optional[Decimal]:
    """Importo come ``Decimal`` (mai float); vuoto o «N/A» = nessun importo."""
    if valore is None or isinstance(valore, bool):
        return None
    if isinstance(valore, (int, Decimal)):
        return Decimal(valore).quantize(Decimal("0.01"))
    if isinstance(valore, float):
        return Decimal(repr(valore)).quantize(Decimal("0.01"))
    testo = re.sub(r"[€\s]", "", str(valore))
    if testo.upper() in _NON_DATO:
        return None
    if "," in testo:  # 1.234,56
        testo = testo.replace(".", "").replace(",", ".")
    try:
        return Decimal(testo).quantize(Decimal("0.01"))
    except InvalidOperation:
        return None


def direzione_da(valore: Any) -> Optional[str]:
    testo = (_testo(valore) or "").lower()
    if testo.startswith("entr"):
        return "entrata"
    if testo.startswith("usc"):
        return "uscita"
    if testo.startswith("intern"):
        return "interno"
    return None


def drive_file_id_da(link: Optional[str]) -> Optional[str]:
    trovato = _DRIVE_ID_RE.search(link or "")
    return trovato.group(1) if trovato else None


def ambito_da(riga: Dict[str, Any]) -> str:
    """Solo un'etichetta informativa: un familiare riconosciuto per nome o CF e
    senza veste aziendale esplicita e' ``personale_familiare``."""
    from app.services.personal_family_registry import (
        is_company_context, is_employment_context, match_family_person,
    )

    testo = " ".join(str(riga.get(c) or "") for c in ("oggetto", "controparte", "nome_file", "note", "pratica"))
    if is_employment_context(testo):
        return AMBITO_AZIENDALE
    if match_family_person(testo) and not is_company_context(testo):
        return AMBITO_FAMILIARE
    return AMBITO_AZIENDALE if is_company_context(testo) else AMBITO_DA_VERIFICARE


CAMPI_IMPRONTA = ("data_protocollo", "tipo_corrispondenza", "tipo_documento", "data_documento", "direzione",
                  "oggetto", "controparte", "pratica", "importo", "note", "nome_file", "sha256", "link")


def impronta_riga(riga: Dict[str, Any]) -> str:
    """Impronta del contenuto descrittivo: due letture uguali danno la stessa."""
    corpo = "\x1f".join(str(riga.get(c) if riga.get(c) is not None else "") for c in CAMPI_IMPRONTA)
    return hashlib.sha256(corpo.encode("utf-8")).hexdigest()


# ── lettura del foglio ───────────────────────────────────────────────────────

def leggi_registro_xlsx(contenuto: bytes) -> Dict[str, Any]:
    """Righe valide e scartate del foglio ``REGISTRO_PROTOCOLLO``.

    Le scartate portano solo numero di riga Excel e motivo (mai il contenuto).
    """
    from openpyxl import load_workbook

    try:
        cartella = load_workbook(io.BytesIO(contenuto), data_only=True)
    except Exception as exc:  # file non xlsx: il messaggio dice solo il tipo
        raise RegistroNonValido("XLSX_NON_LEGGIBILE", "il file non e' un xlsx leggibile",
                                {"errore": type(exc).__name__}) from exc
    nome = next((n for n in cartella.sheetnames if _intestazione(n) == FOGLIO), None)
    if nome is None:
        raise RegistroNonValido("FOGLIO_ASSENTE", f"foglio {FOGLIO} assente", {"fogli": cartella.sheetnames})
    righe_xlsx = cartella[nome].iter_rows()
    try:
        intestazioni = next(righe_xlsx)
    except StopIteration as exc:
        raise RegistroNonValido("FOGLIO_VUOTO", "il foglio non ha intestazioni") from exc
    colonna: Dict[str, int] = {}
    for indice, cella in enumerate(intestazioni):
        chiave = INTESTAZIONI.get(_intestazione(cella.value))
        if chiave and chiave not in colonna:
            colonna[chiave] = indice
    mancanti = [c for c in OBBLIGATORIE if c not in colonna]
    if mancanti:
        raise RegistroNonValido("COLONNE_ASSENTI", "colonne obbligatorie assenti",
                                {"mancanti": mancanti, "trovate": sorted(colonna)})

    valide: Dict[str, Dict[str, Any]] = {}
    scartate: List[Dict[str, Any]] = []
    for n_riga, celle in enumerate(righe_xlsx, start=2):
        if not any(c.value not in (None, "") for c in celle):
            continue

        def valore(chiave: str) -> Any:
            indice = colonna.get(chiave)
            return celle[indice].value if indice is not None and indice < len(celle) else None

        canonico = numero_canonico(valore("numero"))
        if not canonico:
            scartate.append({"riga_excel": n_riga, "motivo": "numero_non_valido"})
            continue
        numero, anno, progressivo = canonico
        sha = (_testo(valore("sha256")) or "").lower() or None
        if sha and not _SHA_RE.match(sha):
            scartate.append({"riga_excel": n_riga, "motivo": "sha256_non_valido", "numero": numero})
            continue
        link = None
        indice_link = colonna.get("link")
        if indice_link is not None and indice_link < len(celle):
            cella = celle[indice_link]
            link = getattr(getattr(cella, "hyperlink", None), "target", None) or _testo(cella.value)
            if link and not str(link).lower().startswith("http"):
                link = None
        data_doc = data_iso(valore("data_documento"))
        oggetto = (_testo(valore("oggetto")) or "")[:MAX_OGGETTO] or None
        importo = importo_decimal(valore("importo"))
        riga = {
            "id": numero, "numero": numero, "anno_protocollo": anno, "progressivo": progressivo,
            "data_protocollo": data_iso(valore("data_protocollo")),
            "tipo_corrispondenza": _testo(valore("tipo_corrispondenza")),
            "tipo_documento": _testo(valore("tipo_documento")),
            "data_documento": data_doc,
            "anno_documento": int(data_doc[:4]) if data_doc else None,
            "direzione": direzione_da(valore("direzione")),
            "oggetto": oggetto, "controparte": _testo(valore("controparte")),
            "pratica": _testo(valore("pratica")),
            "importo": None if importo is None else f"{importo:.2f}",
            "valuta": None if importo is None else "EUR",
            "note": _testo(valore("note")), "nome_file": _testo(valore("nome_file")),
            "sha256": sha, "link": link, "drive_file_id": drive_file_id_da(link),
        }
        precedente = valide.get(numero)
        if precedente is not None:
            uguale = impronta_riga(precedente) == impronta_riga(riga)
            scartate.append({"riga_excel": n_riga, "numero": numero,
                             "motivo": "numero_ripetuto_uguale" if uguale else "numero_ripetuto_diverso"})
            continue
        riga["ambito"] = ambito_da(riga)
        riga["accounting_excluded"] = riga["ambito"] == AMBITO_FAMILIARE
        riga["testo_indice"] = (oggetto or "")[:MAX_TESTO_INDICE] or None
        riga["impronta_riga"] = impronta_riga(riga)
        valide[numero] = riga
    return {"righe": list(valide.values()), "scartate": scartate, "colonne": sorted(colonna)}


# ── import ───────────────────────────────────────────────────────────────────

# Campi che un nuovo import puo' riscrivere: mai stato, ambito, OCR gia' agganciato.
CAMPI_AGGIORNABILI = tuple(c for c in CAMPI_IMPRONTA if c != "sha256") + (
    "anno_documento", "valuta", "drive_file_id", "impronta_riga")


def classifica(righe: Sequence[Dict[str, Any]], esistenti: Dict[str, Dict[str, Any]]) -> Dict[str, Any]:
    """Piano puro: quali righe sono nuove, aggiornate, invariate, in conflitto, rimosse."""
    nuove, aggiornate, invariate, conflitti, rimosse = [], [], [], [], []
    for riga in righe:
        attuale = esistenti.get(riga["id"])
        if attuale is None:
            nuove.append(riga)
        elif attuale.get("stato") == STATO_RIMOSSO:
            rimosse.append(riga["numero"])
        elif attuale.get("sha256") and riga.get("sha256") and attuale["sha256"] != riga["sha256"]:
            conflitti.append(riga["numero"])
        elif attuale.get("impronta_riga") == riga["impronta_riga"]:
            invariate.append(riga["numero"])
        else:
            aggiornate.append(riga)
    nel_file = {r["id"] for r in righe}
    assenti = sorted(i for i, d in esistenti.items() if i not in nel_file and d.get("stato") != STATO_RIMOSSO)
    return {"nuove": nuove, "aggiornate": aggiornate, "invariate": invariate,
            "in_conflitto": conflitti, "rimosse_saltate": rimosse, "assenti_nel_file": assenti}


def _sha_ripetuti(righe: Iterable[Dict[str, Any]]) -> List[str]:
    viste: Dict[str, int] = {}
    for r in righe:
        if r.get("sha256"):
            viste[r["sha256"]] = viste.get(r["sha256"], 0) + 1
    return [s for s, n in viste.items() if n > 1]


async def importa_registro(db, contenuto: bytes, *, dry_run: bool = True,
                           nome_fonte: Optional[str] = None) -> Dict[str, Any]:
    """Importa il registro. Solo contatori e numeri nel risultato (niente dati personali)."""
    letto = leggi_registro_xlsx(contenuto)
    righe = letto["righe"]
    proiezione = {"_id": 0, "id": 1, "stato": 1, "sha256": 1, "impronta_riga": 1}
    esistenti = {d["id"]: d for d in await db[COLL].find({}, proiezione).to_list(None) if d.get("id")}
    piano = classifica(righe, esistenti)
    ora = datetime.now(timezone.utc).isoformat()
    if not dry_run:
        nuove = piano["nuove"]
        for i in range(0, len(nuove), LOTTO_SCRITTURA):
            await db[COLL].insert_many([
                {**r, "stato": STATO_ATTIVO, "canale": CANALE_DRIVE, "fonte_registro": nome_fonte,
                 "importato_il": ora, "aggiornato_il": ora}
                for r in nuove[i:i + LOTTO_SCRITTURA]
            ])
        for riga in piano["aggiornate"]:
            campi = {c: riga.get(c) for c in CAMPI_AGGIORNABILI}
            await db[COLL].update_one({"id": riga["id"]}, {"$set": {**campi, "aggiornato_il": ora}})
    esito = {
        "dry_run": dry_run,
        "righe_nel_file": len(righe),
        "nuovi": len(piano["nuove"]),
        "aggiornati": len(piano["aggiornate"]),
        "invariati": len(piano["invariate"]),
        "in_conflitto": len(piano["in_conflitto"]),
        "rimossi_saltati": len(piano["rimosse_saltate"]),
        "assenti_nel_file": len(piano["assenti_nel_file"]),
        "scartati": len(letto["scartate"]),
        "impronte_ripetute": len(_sha_ripetuti(righe)),
        "elenco_in_conflitto": piano["in_conflitto"][:MAX_ELENCO],
        "elenco_assenti_nel_file": piano["assenti_nel_file"][:MAX_ELENCO],
        "elenco_scartati": letto["scartate"][:MAX_ELENCO],
    }
    logger.info("Protocollo personale import dry_run=%s nuovi=%s aggiornati=%s invariati=%s conflitti=%s scartati=%s",
                dry_run, esito["nuovi"], esito["aggiornati"], esito["invariati"], esito["in_conflitto"],
                esito["scartati"])
    return esito


async def _salva_stato(db, chiave: str, **campi) -> None:
    await db["sistema_stato"].update_one(
        {"chiave": chiave},
        {"$set": {"chiave": chiave, **campi, "updated_at": datetime.now(timezone.utc).isoformat()}},
        upsert=True,
    )


def _chiave(dry_run: bool) -> str:
    return CHIAVE_STATO + (":anteprima" if dry_run else "")


async def _esegui(db, file_id: str, dry_run: bool) -> None:
    chiave = _chiave(dry_run)
    async with _lock:
        await _salva_stato(db, chiave, stato="in_corso", dry_run=dry_run, file_id=file_id,
                           iniziata_il=datetime.now(timezone.utc).isoformat())
        try:
            from app.services.drive_download import scarica_originale

            contenuto = await scarica_originale(file_id)
            esito = await importa_registro(db, contenuto, dry_run=dry_run, nome_fonte=f"drive:{file_id}")
            await _salva_stato(db, chiave, stato="completato", dry_run=dry_run, file_id=file_id, esito=esito)
        except RegistroNonValido as exc:
            await _salva_stato(db, chiave, stato="errore", dry_run=dry_run, file_id=file_id,
                               errore={"code": exc.codice, "message": str(exc), "details": exc.dettagli})
        except Exception as exc:
            logger.exception("Protocollo personale: import fallito (%s)", type(exc).__name__)
            await _salva_stato(db, chiave, stato="errore", dry_run=dry_run, file_id=file_id,
                               errore={"code": "IMPORT_FALLITO", "message": type(exc).__name__})


async def avvia(db, file_id: str, *, dry_run: bool = True) -> Dict[str, Any]:
    """Avvia in background (oltre 5 minuti il proxy Render taglia la richiesta)."""
    global _lavoro
    if _lock.locked() or (_lavoro is not None and not _lavoro.done()):
        return {"avviato": False, **await stato(db)}
    _lavoro = asyncio.create_task(_esegui(db, file_id, dry_run))
    return {"avviato": True, "dry_run": dry_run, "file_id": file_id}


async def stato(db) -> Dict[str, Any]:
    letti: Dict[str, Any] = {}
    for chiave, nome in ((CHIAVE_STATO, "import"), (CHIAVE_STATO + ":anteprima", "anteprima")):
        doc = await db["sistema_stato"].find_one({"chiave": chiave}, {"_id": 0, "chiave": 0})
        if doc:
            letti[nome] = doc
    return letti or {"stato": "mai_avviato"}


# ── rimozione (mai cancellazione) e testo del PDF ────────────────────────────

async def segna_rimosso(db, numero: str, motivo: str, *, attore: str) -> Dict[str, Any]:
    canonico = numero_canonico(numero)
    motivo = " ".join(str(motivo or "").split())
    if not canonico or not motivo:
        return {"success": False, "code": "DATI_NON_VALIDI", "message": "servono un numero valido e il motivo"}
    ora = datetime.now(timezone.utc).isoformat()
    esito = await db[COLL].update_one(
        {"id": canonico[0], "stato": {"$ne": STATO_RIMOSSO}},
        {"$set": {"stato": STATO_RIMOSSO, "rimosso_il": ora, "rimosso_motivo": motivo[:300],
                  "rimosso_da": attore, "aggiornato_il": ora}},
    )
    if not getattr(esito, "matched_count", 0):
        return {"success": False, "code": "NON_TROVATO", "message": "protocollo non trovato o gia' rimosso"}
    return {"success": True, "numero": canonico[0], "stato": STATO_RIMOSSO}


async def ripristina(db, numero: str, *, attore: str) -> Dict[str, Any]:
    canonico = numero_canonico(numero)
    if not canonico:
        return {"success": False, "code": "DATI_NON_VALIDI", "message": "numero non valido"}
    ora = datetime.now(timezone.utc).isoformat()
    esito = await db[COLL].update_one(
        {"id": canonico[0], "stato": STATO_RIMOSSO},
        {"$set": {"stato": STATO_ATTIVO, "ripristinato_il": ora, "ripristinato_da": attore, "aggiornato_il": ora}},
    )
    if not getattr(esito, "matched_count", 0):
        return {"success": False, "code": "NON_TROVATO", "message": "protocollo non trovato o non rimosso"}
    return {"success": True, "numero": canonico[0], "stato": STATO_ATTIVO}


async def allega_testo_pdf(db, numero: str, contenuto_pdf: bytes) -> Dict[str, Any]:
    """Legge il testo di un PDF con il lettore unico (`pdf_text_extraction`) e lo
    aggancia alla riga **solo se l'SHA-256 del file e' quello del registro**: un
    file diverso non si attribuisce a un numero per somiglianza."""
    from app.services.pdf_text_extraction import extract_pdf_text

    canonico = numero_canonico(numero)
    if not canonico:
        return {"success": False, "code": "DATI_NON_VALIDI", "message": "numero non valido"}
    riga = await db[COLL].find_one({"id": canonico[0]}, {"_id": 0, "id": 1, "sha256": 1})
    if not riga:
        return {"success": False, "code": "NON_TROVATO", "message": "protocollo non trovato"}
    if not riga.get("sha256"):
        return {"success": False, "code": "SHA256_ASSENTE",
                "message": "la riga non ha l'impronta del file: non si puo' provare che sia lo stesso"}
    if riga["sha256"] != hashlib.sha256(contenuto_pdf).hexdigest():
        return {"success": False, "code": "IMPRONTA_DIVERSA", "message": "il file non e' quello registrato"}
    testo = ((await asyncio.to_thread(extract_pdf_text, contenuto_pdf)) or "").strip()
    if not testo:
        return {"success": False, "code": "TESTO_ASSENTE",
                "message": "il PDF non ha livello testo: l'OCR sulle scansioni non e' ancora collegato"}
    ora = datetime.now(timezone.utc).isoformat()
    await db[COLL].update_one({"id": canonico[0]}, {"$set": {
        "testo_ocr": testo[:MAX_TESTO_OCR], "testo_indice": testo[:MAX_TESTO_INDICE],
        "testo_fonte": "livello_testo_pdf", "testo_aggiornato_il": ora, "aggiornato_il": ora}})
    return {"success": True, "numero": canonico[0], "caratteri": len(testo)}


# ── ricerca ──────────────────────────────────────────────────────────────────

def normalizza_testo(testo: Any) -> str:
    """Minuscolo senza accenti, **stessa lunghezza** dell'originale (serve per lo snippet)."""
    uscita = []
    for c in str(testo or ""):
        base = unicodedata.normalize("NFD", c)[:1].lower()
        uscita.append(base if len(base) == 1 else " ")
    return "".join(uscita)


def parole_di(q: Any) -> List[str]:
    """Le parole della ricerca: alfanumeriche, senza accenti, senza doppie."""
    viste: List[str] = []
    for parola in re.findall(r"[a-z0-9]+", normalizza_testo(q)):
        if parola not in viste:
            viste.append(parola)
    return viste


CAMPI_TESTO = ("oggetto", "testo_indice", "note", "nome_file", "controparte", "pratica", "tipo_documento",
               "tipo_corrispondenza", "numero", "sha256")


def _giorni(riga: Dict[str, Any]) -> str:
    fuori = []
    for campo in ("data_protocollo", "data_documento"):
        iso = riga.get(campo)
        if iso:
            fuori.append(f"{iso} {iso[8:10]}/{iso[5:7]}/{iso[:4]}")
    return " ".join(fuori)


def _pagliaio(riga: Dict[str, Any]) -> str:
    return normalizza_testo(" ".join([*(str(riga.get(c) or "") for c in CAMPI_TESTO), _giorni(riga)]))


def riga_corrisponde(riga: Dict[str, Any], parole: Sequence[str]) -> bool:
    """AND: ogni parola deve comparire. Una targa o un numero di pratica si
    trovano anche scritti con spazi («GX 037 HJ»)."""
    if not parole:
        return True
    pagliaio = _pagliaio(riga)
    compatto = re.sub(r"[^a-z0-9]", "", pagliaio)
    return all(
        p in pagliaio or (len(p) >= 5 and any(c.isdigit() for c in p) and p in compatto)
        for p in parole
    )


def snippet(riga: Dict[str, Any], parole: Sequence[str], *, raggio: int = 70) -> Optional[List[Dict[str, Any]]]:
    """Frammento di testo attorno alla prima parola trovata, a segmenti
    ``{"t": testo, "hit": bool}``: la pagina evidenzia senza mai costruire HTML."""
    if not parole:
        return None
    for campo in ("testo_indice", "oggetto", "note", "nome_file", "controparte", "pratica"):
        originale = str(riga.get(campo) or "")
        norm = normalizza_testo(originale)
        posizioni = sorted((m.start(), m.end()) for p in parole for m in re.finditer(re.escape(p), norm))
        if not posizioni:
            continue
        inizio = max(0, posizioni[0][0] - raggio)
        fine = min(len(originale), posizioni[0][1] + raggio)
        segmenti: List[Dict[str, Any]] = []
        cursore = inizio
        for a, b in posizioni:
            if a < cursore or b > fine:
                continue
            if a > cursore:
                segmenti.append({"t": originale[cursore:a], "hit": False})
            segmenti.append({"t": originale[a:b], "hit": True})
            cursore = b
        if cursore < fine:
            segmenti.append({"t": originale[cursore:fine], "hit": False})
        if inizio > 0:
            segmenti.insert(0, {"t": "…", "hit": False})
        if fine < len(originale):
            segmenti.append({"t": "…", "hit": False})
        return segmenti
    return None


def _chiave_ordine(riga: Dict[str, Any]) -> Tuple[str, str]:
    return (riga.get("data_documento") or riga.get("data_protocollo") or "", riga.get("numero") or "")


def filtra(righe: Iterable[Dict[str, Any]], *, q: Optional[str] = None, anno: Optional[int] = None,
           tipo_documento: Optional[str] = None, ambito: Optional[str] = None,
           includi_rimossi: bool = False) -> List[Dict[str, Any]]:
    """Ricerca AND sulle parole, filtri esatti, piu' recenti per prime (puro, testabile).

    ``anno`` vale per l'anno del documento **o** del protocollo."""
    parole = parole_di(q)
    tipo = normalizza_testo(tipo_documento).strip() if tipo_documento else ""
    trovate = []
    for riga in righe:
        if not includi_rimossi and riga.get("stato") == STATO_RIMOSSO:
            continue
        if anno is not None and anno not in (riga.get("anno_documento"), riga.get("anno_protocollo")):
            continue
        if tipo and tipo not in normalizza_testo(riga.get("tipo_documento")):
            continue
        if ambito and riga.get("ambito") != ambito:
            continue
        if riga_corrisponde(riga, parole):
            trovate.append(riga)
    trovate.sort(key=_chiave_ordine, reverse=True)
    return trovate


def _pubblica(riga: Dict[str, Any], parole: Sequence[str]) -> Dict[str, Any]:
    fuori = {k: v for k, v in riga.items() if k not in ("testo_ocr", "testo_indice", "_id", "impronta_riga")}
    fuori["snippet"] = snippet(riga, parole)
    return fuori


async def cerca(db, *, q: Optional[str] = None, anno: Optional[int] = None, tipo_documento: Optional[str] = None,
                ambito: Optional[str] = None, includi_rimossi: bool = False,
                limit: int = 50, offset: int = 0) -> Dict[str, Any]:
    """Elenco senza payload (`testo_ocr` escluso), a pagine di al piu' 200 righe."""
    limit = max(1, min(int(limit), LIMITE_PAGINA))
    offset = max(0, int(offset))
    righe = await db[COLL].find({}, metadata_projection(COLL)).to_list(None)
    trovate = filtra(righe, q=q, anno=anno, tipo_documento=tipo_documento, ambito=ambito,
                     includi_rimossi=includi_rimossi)
    parole = parole_di(q)
    return {
        "totale": len(trovate), "limit": limit, "offset": offset, "parole": parole,
        "righe": [_pubblica(r, parole) for r in trovate[offset:offset + limit]],
    }


# ── ponte informativo verso la contabilita' ─────────────────────────────────

# (collezione, campo con lo SHA-256 del PDF, etichetta): sola lettura
SORGENTI_IMPRONTA: Tuple[Tuple[str, str, str], ...] = (
    ("cartelle_pagamento", "sha256", "cartella_pagamento"),
    ("atti_giudiziari", "id", "atto_giudiziario"),
    ("pagopa_receipts", "pdf_hash", "ricevuta_pagopa"),
    ("quietanze_f24", "pdf_hash", "quietanza_f24"),
    ("f24_unificato", "pdf_hash", "modello_f24"),
    ("bonifici_transfers", "pdf_hash", "bonifico"),
    ("documents_inbox", "sha256", "documento_caricato"),
)


async def documenti_collegati(db, riga: Dict[str, Any]) -> Dict[str, Any]:
    """Cosa la contabilita' sa gia' di questo documento. **Solo lettura**: non
    crea relazioni, scritture o pagamenti, e non cambia lo stato di niente."""
    trovati: List[Dict[str, Any]] = []
    impronta = riga.get("sha256")
    if impronta:
        for collezione, campo, etichetta in SORGENTI_IMPRONTA:
            try:
                documenti = await db[collezione].find({campo: impronta}, {"_id": 0, "id": 1}).to_list(5)
            except Exception as exc:
                logger.warning("Ponte protocollo: lettura di %s non riuscita (%s)", collezione, type(exc).__name__)
                continue
            for documento in documenti:
                trovati.append({"tipo": etichetta, "collezione": collezione,
                                "id": documento.get("id") or impronta, "via": "impronta_sha256"})
    drive_id = riga.get("drive_file_id")
    if drive_id:
        try:
            from app.services.entity_relations import find_entity_relations

            for rel in await find_entity_relations(db, entity_type="documento", entity_id=drive_id, limit=50):
                if rel.get("status") == "revoked":
                    continue
                origine = rel.get("source") or {}
                altro = rel.get("target") if origine.get("type") == "documento" and origine.get("id") == drive_id \
                    else origine
                trovati.append({"tipo": (altro or {}).get("type"), "collezione": None, "id": (altro or {}).get("id"),
                                "via": "entity_relations", "stato": rel.get("status")})
        except Exception as exc:
            logger.warning("Ponte protocollo: lettura delle relazioni non riuscita (%s)", type(exc).__name__)
    return {
        "sola_lettura": True,
        "escluso_dalla_contabilita": bool(riga.get("accounting_excluded")),
        "documenti": trovati,
    }


async def dettaglio(db, numero: str) -> Optional[Dict[str, Any]]:
    canonico = numero_canonico(numero)
    if not canonico:
        return None
    riga = await db[COLL].find_one({"id": canonico[0]}, {"_id": 0})
    if not riga:
        return None
    return {**{k: v for k, v in riga.items() if k != "impronta_riga"},
            "collegati": await documenti_collegati(db, riga)}
