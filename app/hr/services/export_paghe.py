"""Proiezione Excel della vista paghe: documenti e pagamenti restano distinti."""
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MESI = ['Gennaio', 'Febbraio', 'Marzo', 'Aprile', 'Maggio', 'Giugno', 'Luglio',
        'Agosto', 'Settembre', 'Ottobre', 'Novembre', 'Dicembre', 'Tredicesima', 'Quattordicesima']


def formato_fogli(wb):
    for ws in wb:
        ws.freeze_panes = 'C2'
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='426054')
            cell.alignment = Alignment(wrap_text=True, vertical='center')
        ws.row_dimensions[1].height = 32
        for cells in ws.iter_rows(min_row=2):
            for cell in cells:
                # Testo proveniente da documenti: non deve diventare una formula.
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                if isinstance(cell.value, (float, Decimal)):
                    cell.number_format = '#,##0.00 "€"'
                if hasattr(cell.value, 'year'):
                    cell.number_format = 'dd/mm/yyyy'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        for col in ws.columns:
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(48, max(14, max(len(str(c.value or '')) for c in col) + 2))


def workbook_paghe(dati):
    wb = Workbook()
    ced = wb.active
    ced.title = 'Importi cedolini'
    ced.append(['Dipendente', 'Anno', 'Mese', 'Netto estratto PDF EUR', 'Netto confermato da elenco EUR',
                'Recupero acconto in busta EUR', 'Dovuto EUR', 'PDF presente', 'ID cedolino', 'Da verificare'])
    pag = wb.create_sheet('Pagamenti recuperati')
    pag.append(['Dipendente', 'Anno competenza', 'Mese competenza', 'Data pagamento', 'Importo EUR',
                'Tipo', 'Riferimento', 'Fonte', 'Associazione confermata', 'Nota'])
    missing = wb.create_sheet('PDF non collegati')
    missing.append(['Dipendente', 'Anno', 'Mese', 'Importo noto EUR', 'Da verificare'])
    rec = wb.create_sheet('Mesi riconciliati')
    rec.append(['Dipendente', 'Anno', 'Mese', 'Dovuto EUR', 'Pagamenti EUR', 'Residuo EUR', 'Fonte', 'Riferimenti'])
    keys = set()
    for r in dati['righe']:
        key = (r.get('dipendente_id'), r.get('anno'), r.get('mese'))
        keys.add(key)
        mese = MESI[int(r['mese']) - 1] if r.get('mese') and 1 <= int(r['mese']) <= 14 else str(r.get('mese') or '')
        ident = [r.get('dipendente'), r.get('anno'), mese]
        avvisi = []
        if r.get('avvisi_importo'):
            avvisi.append('Importi discordanti: confrontare le fonti nella pagina paghe')
        if r.get('acconto_da_verificare'):
            avvisi.append('Recupero acconto da verificare')
        if r.get('busta') is None:
            avvisi.append('Importo cedolino mancante')
        if r.get('cedolino_pdf') and r.get('netto_stampato') is None:
            avvisi.append('Netto PDF da leggere')
        ced.append(ident + [r.get('netto_stampato'), r.get('netto_confermato'), r.get('acconto_recuperato'),
                           r.get('busta'), 'Sì' if r.get('cedolino_pdf') else 'No', r.get('cedolino_id'), '; '.join(avvisi)])
        if not r.get('cedolino_pdf'):
            missing.append(ident + [r.get('busta'), 'Nessun PDF collegato al periodo; disponibilità da verificare. Il pagamento resta nel saldo.'])
        bonifici = r.get('bonifici') or []
        riferimenti = ', '.join(dict.fromkeys(str(b['riferimento']) for b in bonifici if b.get('riferimento')))
        for b in bonifici:
            pag.append(ident + [b.get('data'), b.get('importo'), 'Bonifico', b.get('riferimento'), r.get('fonte'),
                               'Sì' if r.get('riconciliato') else 'No', b.get('causale')])
        if not bonifici and r.get('bonifico'):
            pag.append(ident + [r.get('bonifico_data'), r['bonifico'], 'Importo aggregato', '', r.get('fonte'),
                               'Sì' if r.get('riconciliato') else 'No', 'Dettaglio del bonifico non presente in questa vista'])
        for a in r.get('acconti_dettaglio') or []:
            pag.append(ident + [a.get('data'), a.get('importo'), 'Acconto ammesso', '', 'Registro acconti',
                               'Sì' if r.get('riconciliato') else 'No', 'Distinto dal recupero acconto trattenuto nel cedolino'])
        if r.get('riconciliato') and r.get('stato') == 'pagato' and r.get('fonte') == 'banca' and r.get('cedolino_pdf') and (r.get('erogato') or 0) > 0 and not avvisi:
            rec.append(ident + [r.get('busta'), r.get('erogato'), r.get('saldo'), r.get('fonte'), riferimenti])
    for r in dati.get('cedolini_da_verificare') or []:
        if (r.get('dipendente_id'), r.get('anno'), r.get('mese')) in keys:
            continue
        mese = MESI[int(r['mese']) - 1] if r.get('mese') and 1 <= int(r['mese']) <= 14 else str(r.get('mese') or '')
        ced.append([r.get('dipendente'), r.get('anno'), mese, None, None, None, None, 'Sì', r.get('cedolino_id'), 'Netto PDF da leggere'])
    formato_fogli(wb)
    return wb
