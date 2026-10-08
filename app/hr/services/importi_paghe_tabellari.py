"""Importi tabellari: netto dichiarato distinto dalle ricostruzioni dei pagamenti."""
from __future__ import annotations

import csv
import asyncio
import hashlib
import io
import json
import re
import unicodedata
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from zipfile import BadZipFile

from app.services.cedolini_motore import PAYROLL_MIN_YEAR

MESI = {nome: i + 1 for i, nome in enumerate((
    "gennaio", "febbraio", "marzo", "aprile", "maggio", "giugno", "luglio",
    "agosto", "settembre", "ottobre", "novembre", "dicembre", "tredicesima", "quattordicesima"))}


def normalizza(value):
    return re.sub(r"[^a-z0-9]", "", unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode().lower())


def euro(value, *, allow_negative=False):
    if value is None or str(value).strip() == "":
        return None
    text = str(value).replace("€", "").replace("EUR", "").replace("\u00a0", "").replace(" ", "").strip()
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        number = Decimal(text)
        if not number.is_finite() or (number < 0 and not allow_negative) or number != number.quantize(Decimal(".01")):
            raise ValueError("importo negativo, non finito o non al centesimo")
        return number
    except InvalidOperation as exc:
        raise ValueError("importo non valido") from exc


def intero(value):
    try:
        number = Decimal(str(value))
        if not number.is_finite() or number != number.to_integral_value():
            raise ValueError("numero intero richiesto")
        return int(number)
    except (InvalidOperation, TypeError) as exc:
        raise ValueError("numero intero richiesto") from exc


def leggi(content: bytes):
    if len(content) > 10 * 1024 * 1024:
        raise ValueError("File oltre 10 MB")
    if content.startswith(b"PK"):
        import openpyxl
        try:
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
        except (BadZipFile, KeyError, OSError, ValueError) as exc:
            raise ValueError("File Excel non valido: carica un .xlsx leggibile") from exc
        try:
            # Gli estratti della chat contengono prima «Pagamenti», poi
            # «Cedolini»: i bonifici singoli non sono importi netti di busta.
            nome_foglio = next((s for nome in ("salari", "cedolini") for s in wb.sheetnames if normalizza(s) == nome), wb.sheetnames[0])
            ws = wb[nome_foglio]
            if ws.max_row > 10001:
                raise ValueError("Massimo 10.000 righe per import")
            rows = list(ws.iter_rows(values_only=True))
        finally:
            wb.close()
    else:
        text = content.decode("utf-8-sig")
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters="\t;,")
        except csv.Error as exc:
            raise ValueError("Usa un Excel, CSV o testo con colonne separate da tabulazioni") from exc
        rows = list(csv.reader(io.StringIO(text), dialect))
    if not rows or len(rows) > 10001:
        raise ValueError("Foglio vuoto o oltre 10.000 righe")
    header = [normalizza(v) for v in rows[0]]

    def column(*names):
        return next((i for i, name in enumerate(header) if name in {normalizza(n) for n in names}), None)

    cols = {
        "nome": column("Dipendente", "Nome dipendente"), "cf": column("Codice fiscale", "CF"),
        "matricola": column("Matricola"), "iban": column("IBAN"),
        "competenza": column("Competenza"), "mese": column("Mese", "Mese cedolino"), "anno": column("Anno"),
        "netto": column("Stipendio Netto", "Netto", "Importo busta", "Netto cedolino"),
        "attribuito": column("Importo attribuito EUR", "Importo attribuito"),
        "erogato": column("Importo Erogato", "Erogato", "Bonifico", "Uscita"),
        "ricostruzione": column("Tipo ricostruzione"), "stato": column("Stato"),
    }
    if cols["nome"] is None or (cols["competenza"] is None and (cols["mese"] is None or cols["anno"] is None)):
        raise ValueError("Servono Dipendente e Competenza (oppure Mese e Anno)")
    if all(cols[k] is None for k in ("netto", "attribuito", "erogato")):
        raise ValueError("Manca una colonna Netto, Importo busta, Importo attribuito EUR o Importo Erogato")
    output, errors = [], []
    for rownum, row in enumerate(rows[1:], 2):
        def cell(k):
            i = cols[k]
            return row[i] if i is not None and i < len(row) else None
        if not any(str(v or "").strip() for v in row):
            continue
        try:
            if not cell("nome"):
                raise ValueError("dipendente mancante")
            period = str(cell("competenza") or "").strip().lower()
            if isinstance(cell("competenza"), datetime):
                year, month = cell("competenza").year, cell("competenza").month
            elif period:
                match = re.fullmatch(r"([a-zà]+|\d{1,2})[ /-]+(\d{4})", period)
                reverse = re.fullmatch(r"(\d{4})-(\d{1,2})", period)
                if match:
                    m, y = match.groups()
                    year, month = int(y), MESI.get(m, int(m) if m.isdigit() else 0)
                elif reverse:
                    year, month = map(int, reverse.groups())
                else:
                    raise ValueError("competenza non valida")
            else:
                m = normalizza(cell("mese"))
                year = intero(cell("anno"))
                month = MESI[m] if m in MESI else intero(cell("mese"))
            if not PAYROLL_MIN_YEAR <= year <= datetime.now().year + 1 or not 1 <= month <= 14:
                raise ValueError("mese o anno fuori intervallo")
            values = {k: euro(cell(k)) for k in ("netto", "attribuito", "erogato")}
            kind = next((k for k in ("netto", "attribuito", "erogato") if values[k] is not None), None)
            if not kind:
                raise ValueError("importo mancante")
            matricola = cell("matricola")
            if isinstance(matricola, (int, float)):
                matricola = str(intero(matricola))
            output.append({"riga": rownum, "nome": str(cell("nome")).strip(), "anno": year, "mese": month,
                           "tipo": kind, "importo": float(values[kind]),
                           "cf": normalizza(cell("cf")), "matricola": normalizza(matricola).lstrip("0"),
                           "iban": normalizza(cell("iban")), "nota": " · ".join(str(cell(k)) for k in ("ricostruzione", "stato") if cell(k))})
        except (ValueError, TypeError) as exc:
            errors.append(f"Riga {rownum}: {exc}")
    return output, errors


def avvisi(paga):
    if not paga.get("importi_excel"):
        return []
    out = []
    for entry in paga.get("importi_excel") or []:
        # Confrontare un netto con il netto del PDF, una ricostruzione dei
        # pagamenti con il dovuto (netto + recupero): grandezze distinte.
        current = euro(paga.get("netto_stampato") if entry.get("tipo") == "netto" and paga.get("netto_stampato") is not None and (not paga.get("importo_busta_manuale") or paga.get("netto_confermato") is not None) else paga.get("importo_busta"), allow_negative=True)
        expected = euro(entry.get("importo"))
        conflict = current is not None and expected != current
        if not conflict and current is not None and entry.get("tipo") == "netto":
            continue
        if entry.get("verificato") and entry.get("verificato_importo") == paga.get("importo_busta") and (entry.get("tipo") != "netto" or entry.get("verificato_netto") == paga.get("netto_stampato")):
            continue
        out.append({**entry, "busta_attuale": float(current) if current is not None else None,
                    "differenza": float(expected - current) if current is not None else None,
                    "motivo": "Importi discordanti nello stesso file: verificare il periodo" if entry.get("conflitto_file") else "Importo diverso dal cedolino" if conflict else "Netto da verificare" if entry.get("tipo") == "netto" else "Importo ricostruito dai pagamenti: verificare il cedolino"})
    return out


async def importa(db, content: bytes, filename: str, *, applica=False):
    from app.hr.db_supabase import SupabaseDatabase
    from app.hr.services.sincronizza_paghe_mensili import _SYNC_LOCK
    rows, errors = await asyncio.to_thread(leggi, content)
    result = {"aggiornati": 0, "duplicati": 0, "non_trovati": 0, "nomi_non_trovati": [],
              "righe_aggregate": len(rows), "errori": errors, "discrepanze": [], "da_verificare": 0,
              "anteprima": not applica, "mesi": [], "success": True}
    file_hash = hashlib.sha256(content).hexdigest()
    async with _SYNC_LOCK:
        if isinstance(db, SupabaseDatabase):
            await db.refresh_collections("dipendenti", "paghe_mensili")
        employees = await db.dipendenti.find({"merged_into": {"$exists": False}}, {"_id": 0}).to_list(None)
        indexes = {k: {} for k in ("nome", "cf", "matricola", "iban")}
        for d in employees:
            names = [d.get("nome_completo"), f"{d.get('cognome','')} {d.get('nome','')}", f"{d.get('nome','')} {d.get('cognome','')}"]
            for k, vals in {"nome": names, "cf": [d.get("codice_fiscale")], "matricola": [d.get("matricola")], "iban": [d.get("iban")]}.items():
                for value in vals:
                    key = normalizza(value)
                    if k == "matricola": key = key.lstrip("0")
                    if key: indexes[k].setdefault(key, set()).add(d["id"])
        existing = {(p.get("dipendente_id"), p.get("anno"), p.get("mese")): p
                    async for p in db.paghe_mensili.find({}, {"_id": 0})}
        resolved = []
        netti_file = {}
        for row in rows:
            matches = indexes["nome"].get(normalizza(row["nome"]), set())
            if row["cf"]:
                matches = matches & indexes["cf"].get(row["cf"], set())
            for field in ("matricola", "iban"):
                candidates = indexes[field].get(row[field], set())
                if candidates: matches = matches & candidates
            if len(matches) != 1:
                result["non_trovati"] += 1
                result["nomi_non_trovati"].append(row["nome"])
                continue
            dip = next(iter(matches))
            resolved.append((row, dip))
            if row["tipo"] == "netto":
                netti_file.setdefault((dip, row["anno"], row["mese"]), set()).add(euro(row["importo"]))
        seen = set()
        for row, dip in resolved:
            key = (dip, row["anno"], row["mese"])
            old = existing.get(key, {})
            identity = hashlib.sha256(json.dumps([*key, row["tipo"], str(euro(row["importo"]))]).encode()).hexdigest()
            entries = list(old.get("importi_excel") or [])
            netto_attuale = old.get("netto_stampato") if old.get("netto_stampato") is not None and not old.get("importo_busta_manuale") else old.get("importo_busta")
            if identity in seen or any(e.get("id") == identity for e in entries) or (row["tipo"] == "netto" and euro(netto_attuale, allow_negative=True) == euro(row["importo"])):
                result["duplicati"] += 1
                continue
            seen.add(identity)
            entry = {k: row[k] for k in ("importo", "tipo", "nota", "riga")}
            if len(netti_file.get(key, ())) > 1:
                entry["conflitto_file"] = True
            entry.update(id=identity, file=filename, sha256=file_hash)
            entries.append(entry)
            patch = {"dipendente_id": dip, "anno": row["anno"], "mese": row["mese"], "importi_excel": entries}
            if row["tipo"] == "netto" and old.get("importo_busta") is None and not entry.get("conflitto_file"):
                patch.update(importo_busta=row["importo"], origine="excel_cedolino", importo_busta_manuale=True,
                             netto_confermato=row["importo"],
                             importo_busta_nota=f"Import Excel: {filename}, riga {row['riga']}")
            updated = {**old, **patch}
            alerts = avvisi(updated)
            if any(a["id"] == identity for a in alerts):
                result["da_verificare"] += 1
                result["discrepanze"].append({"dipendente": row["nome"], "mese": row["mese"], "anno": row["anno"],
                                             "busta_app": old.get("importo_busta"), "busta_excel": row["importo"], "tipo": row["tipo"]})
            if applica:
                patch["updated_at"] = datetime.now(timezone.utc).isoformat()
                await db.paghe_mensili.update_one({"dipendente_id": dip, "anno": row["anno"], "mese": row["mese"]}, {"$set": patch}, upsert=True)
            existing[key] = updated
            result["aggiornati"] += 1
            result["mesi"].append({"anno": row["anno"], "mese": row["mese"]})
    result["nomi_non_trovati"] = sorted(set(result["nomi_non_trovati"]))
    result["mesi"] = [dict(anno=a, mese=m) for a,m in sorted({(p['anno'],p['mese']) for p in result['mesi']})]
    result["imported"] = result["aggiornati"] if applica else 0
    result["duplicates"] = result["duplicati"]
    result["partial"] = bool(errors or result["non_trovati"])
    result["success"] = not result["partial"]
    return result


async def elabora_file(filename: str, content: bytes):
    from app.hr.database import Database
    return await importa(Database.get_db(), content, filename, applica=True)
